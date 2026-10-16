#!/usr/bin/env python3
"""
Tests for TAXII Server Implementation
Tests all TAXII 2 endpoints, ingestion (replace + merge modes), auth,
Vision One integration, and the OTX community-intel puller.
"""

import sys
import os
import base64
import io
import re
import tempfile
import time
import types
import unittest
import urllib.parse
from datetime import datetime, timedelta

import jwt
import sqlalchemy

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

# Tests always run on an isolated SQLite file DB — never pick up DATABASE_URL
# from a local .env (e.g. a dev checkout pointed at PostgreSQL). The .env
# loader only sets vars that are not already in the process environment, so
# setting this before importing server pins the test DB.
os.environ['DATABASE_URL'] = 'sqlite:///' + os.path.join(
    tempfile.gettempdir(), 'taxii2-tests.db')
# Also pin the Get Objects response shape to the spec envelope so a local
# .env (e.g. TAXII_OBJECTS_SHAPE=bundle for a Vision One experiment) can't
# change what the suite asserts.
os.environ['TAXII_OBJECTS_SHAPE'] = 'envelope'

from server import (
    app,
    ThreatIntel,
    generate_taxii_bundle,
    ingest_objects,
    memory_store,
    init_db,
    OtxPoller,
    VisionOnePoller,
    STIXObject,
    sweep_expired_indicators,
)
import server as _server


def _reset_db():
    """Purge feed rows, API keys and dashboard users between tests."""
    import server
    session = server.create_session()
    try:
        memory_store.clear()
        session.query(server.STIXObject).delete()
        session.query(server.ApiKey).delete()
        session.query(server.UIUser).delete()
        session.query(server.AuditEvent).delete()
        session.query(server.BenignValue).delete()
        session.commit()
    finally:
        session.close()
    # Drop rows AND the in-memory benign cache so a test can't inherit an entry
    # a previous test created.
    server.load_benign_values()


def _manual_obj(value="192.168.1.100"):
    return {
        'id': f"ipv4-addr--{value.replace('.', '-')}",
        'type': 'ipv4-addr',
        'object': {'ipv4-addr': {'value': value}},
        'labels': ['threat'],
        'confidence': 90,
    }


def _otx_obj(value="8.8.8.8"):
    return {
        'id': f"ipv4-addr--{value.replace('.', '-')}",
        'type': 'ipv4-addr',
        'object': {'ipv4-addr': {'value': value}},
        'labels': ['otx', 'some-pulse'],
        'confidence': 70,
    }


def _row(last_seen=None, modified=None):
    """A minimal fake STIXObject row carrying the timestamps the gate uses."""
    r = types.SimpleNamespace()
    r.last_seen = last_seen
    r.modified = modified
    return r


# Served objects are STIX 2.1 *indicators* (pattern + valid_from); these read
# the underlying IOC value/kind back out of the pattern so tests assert on the
# intel rather than the wire type.
_PATTERN_VALUE_RE = re.compile(r"= '([^']+)'")


def served_value(o):
    """The IOC value a served STIX object carries (indicator pattern / SCO value)."""
    if isinstance(o, dict) and o.get('value'):
        return o['value']
    m = _PATTERN_VALUE_RE.search((o or {}).get('pattern') or '')
    return m.group(1) if m else None


def served_kind(o):
    """The underlying IOC kind of a served object (derived from its pattern)."""
    p = (o or {}).get('pattern') or ''
    for kind, prefix in (
        ('ipv4-addr', '[ipv4-addr:'), ('domain-name', '[domain-name:'),
        ('file-hash', '[file:'), ('url', '[url:'), ('email-addr', '[email-addr:'),
        ('ipv6-addr', '[ipv6-addr:'), ('mac-addr', '[mac-addr:'),
        ('autonomous-system', '[autonomous-system:'),
        ('windows-registry-key', '[windows-registry-key:'),
    ):
        if p.startswith(prefix):
            return kind
    return (o or {}).get('type')


class TestTaxiiServer(unittest.TestCase):
    """Test cases for TAXII server endpoints"""

    def setUp(self):
        """Set up test environment"""
        self.client = app.test_client()
        # Reset in-memory store + database so tests are isolated
        _reset_db()

        # Configure test TAXII auth credentials
        app.config['TAXII_AUTH'] = {
            'username': 'test_taxii_user',
            'password': 'test_taxii_pass',
        }
        # Feed-wide writes (purge-all / delete / revoke / community pull) need
        # an admin principal; stand one in for the legacy global override.
        app.config['TAXII_ADMIN_AUTH'] = {
            'username': 'test_taxii_user',
            'password': 'test_taxii_pass',
        }
        self.test_auth_headers = {
            'X-Taxii-Username': 'test_taxii_user',
            'X-Taxii-Password': 'test_taxii_pass',
        }

        # Configure test UI credentials (independent of the live .env).
        self.ui_user = 'test_ui_user'
        self.ui_pass = 'test_ui_pass'
        import server
        server._ui_username = self.ui_user
        server._ui_password = self.ui_pass

    def _get_auth_headers(self):
        """Get auth headers for TAXII 2 requests (valid for data endpoints)."""
        return self.test_auth_headers

    def _ingest(self, stix_objects):
        """POST /feed/ingest with auth headers."""
        return self.client.post(
            '/feed/ingest',
            json={'stix_objects': stix_objects},
            headers=self._get_auth_headers(),
        )

    def test_health_check(self):
        """Test health check endpoint"""
        response = self.client.get('/health')
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertEqual(data['status'], 'healthy')
        self.assertIn('objects_count', data)
        self.assertEqual(data['objects_count'], 0)

    def test_get_feed_empty(self):
        """Test feed endpoint with no objects (with auth)"""
        response = self.client.get('/feed', headers=self._get_auth_headers())
        self.assertEqual(response.status_code, 200)
        self.assertIn('Response', response.data.decode('utf-8'))

    def test_get_feed_no_auth(self):
        """Test feed endpoint without auth returns 401"""
        response = self.client.get('/feed')
        self.assertEqual(response.status_code, 401)

    def test_ingest_ip_address(self):
        """Test ingestion of IP address object"""
        response = self._ingest([
            {
                'id': 'ipv4-addr--123',
                'type': 'ipv4-addr',
                'object': {'ipv4-addr': {'value': '192.168.1.100'}},
                'labels': ['threat', 'malicious'],
                'confidence': 80,
            }
        ])
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertEqual(data['message'], 'Data ingested successfully')
        self.assertEqual(data['objects_count'], 1)

    def test_ingest_no_auth(self):
        """Test ingestion without auth returns 401"""
        response = self.client.post(
            '/feed/ingest',
            json={'stix_objects': [
                {'id': 'ipv4-addr--1', 'type': 'ipv4-addr',
                 'object': {'ipv4-addr': {'value': '1.2.3.4'}}}
            ]},
        )
        self.assertEqual(response.status_code, 401)

    def test_ingest_file_hash(self):
        """Test ingestion of file hash object"""
        response = self._ingest([
            {
                'id': 'file-hash--abcdef1234567890',
                'type': 'file-hash',
                'object': {
                    'hash_value': {
                        'algorithm': 'sha256',
                        'value': 'a' * 64,
                    }
                },
                'labels': ['sha256', 'malware'],
                'confidence': 75,
            }
        ])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()['objects_count'], 1)

    def test_ingest_domain_name(self):
        """Test ingestion of domain name object"""
        response = self._ingest([
            {
                'id': 'domain-name--789',
                'type': 'domain-name',
                'object': {'domain-name': {'value': 'malicious.example.com'}},
                'labels': ['threat', 'malicious'],
                'confidence': 65,
            }
        ])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()['objects_count'], 1)

    def test_ingest_indicator(self):
        """Test ingestion of indicator object"""
        response = self._ingest([
            {
                'id': 'indicator--abc',
                'type': 'indicator',
                'object': {'indicator': {'value': 'Suspicious Activity', 'labels': ['threat']}},
                'confidence': 90,
            }
        ])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()['objects_count'], 1)

    def test_get_feed_after_ingest(self):
        """Test feed retrieval after ingestion (with auth)"""
        self._ingest([
            {
                'id': 'ipv4-addr--123',
                'type': 'ipv4-addr',
                'object': {'ipv4-addr': {'value': '192.168.1.100'}},
            }
        ])
        response = self.client.get('/feed', headers=self._get_auth_headers())
        self.assertEqual(response.status_code, 200)
        xml = response.data.decode('utf-8')
        self.assertIn('192.168.1.100', xml)

    def test_purge_data(self):
        """Test data purge endpoint"""
        self._ingest([
            {
                'id': 'ipv4-addr--123',
                'type': 'ipv4-addr',
                'object': {'ipv4-addr': {'value': '192.168.1.100'}},
            }
        ])
        response = self.client.delete('/feed/purge', headers=self._get_auth_headers())
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertEqual(data['message'], 'All data purged successfully')

        # Verify store is empty
        self.assertEqual(len(memory_store), 0)

    def test_purge_no_auth(self):
        """Test purge without auth returns 401"""
        response = self.client.delete('/feed/purge')
        self.assertEqual(response.status_code, 401)

    # ------------------------------------------------------------------
    # Web UI session auth
    # ------------------------------------------------------------------
    def test_ui_session_unauthenticated(self):
        """Test /ui/session reports not authenticated without a cookie."""
        response = self.client.get('/ui/session')
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertFalse(data['authenticated'])
        self.assertIsNone(data['user'])

    def test_ui_login_success_sets_cookie(self):
        """Test /ui/login with correct ui.auth creds sets a session cookie."""
        # UI creds are set to known test values in setUp (independent of .env).
        response = self.client.post(
            '/ui/login',
            json={'username': self.ui_user, 'password': self.ui_pass},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.set_cookie)  # a session cookie was set
        # The session is now valid.
        session = self.client.get('/ui/session').get_json()
        self.assertTrue(session['authenticated'])
        self.assertEqual(session['user'], self.ui_user)

    def test_ui_login_wrong_password(self):
        """Test /ui/login rejects a wrong password."""
        response = self.client.post(
            '/ui/login',
            json={'username': self.ui_user, 'password': 'wrong-password-xyz'},
        )
        self.assertEqual(response.status_code, 401)

    def test_ui_login_unconfigured_fails_closed(self):
        """With no UI creds configured, /ui/login must reject (503), not
        accept empty==empty credentials."""
        import server
        saved = (server._ui_username, server._ui_password)
        try:
            server._ui_username = ''
            server._ui_password = ''
            response = self.client.post(
                '/ui/login', json={'username': '', 'password': ''},
            )
            self.assertEqual(response.status_code, 503)
            self.assertIn('not configured', response.get_json()['error'])
        finally:
            server._ui_username, server._ui_password = saved

    def test_ui_session_cookie_grants_data_access(self):
        """Test a valid UI session cookie can read /objects (no TAXII headers)."""
        import server
        user, pw = self.ui_user, self.ui_pass
        login = self.client.post('/ui/login', json={'username': user, 'password': pw})
        cookie = None
        for k, v in login.headers:
            if k.lower() == 'set-cookie' and 'taxii2_ui_session' in v:
                cookie = v.split('=', 1)[1].split(';', 1)[0]
        self.assertIsNotNone(cookie, "login should set a session cookie")
        # Use the cookie directly (no X-Taxii-* headers) to read /objects.
        response = self.client.get(
            '/objects',
            headers={'Cookie': f'taxii2_ui_session={cookie}'},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn('count', response.get_json())

    def test_ui_logout_clears_cookie(self):
        """Test /ui/logout clears the session cookie."""
        self.client.post('/ui/login', json={'username': self.ui_user, 'password': self.ui_pass})
        response = self.client.post('/ui/logout')
        self.assertEqual(response.status_code, 200)
        set_cookies = [v for k, v in response.headers if k.lower() == 'set-cookie']
        self.assertTrue(any('taxii2_ui_session=' in c for c in set_cookies))
        # After logout the session no longer authenticates.
        self.assertFalse(self.client.get('/ui/session').get_json()['authenticated'])

    def test_objects_accepts_session_or_taxii(self):
        """Test /objects accepts EITHER a session cookie OR TAXII headers."""
        import server
        # With TAXII headers (no cookie).
        r1 = self.client.get('/objects', headers=self._get_auth_headers())
        self.assertEqual(r1.status_code, 200)
        # With a session cookie (no headers).
        self.client.post('/ui/login', json={'username': self.ui_user, 'password': self.ui_pass})
        r2 = self.client.get('/objects')
        self.assertEqual(r2.status_code, 200)

    def test_objects_rejects_no_creds(self):
        """Test /objects with neither cookie nor TAXII headers returns 401."""
        self.client.delete_cookie('taxii2_ui_session')
        response = self.client.get('/objects')
        self.assertEqual(response.status_code, 401)

    # ------------------------------------------------------------------
    # Third-party TAXII 2.1 puller
    # ------------------------------------------------------------------
    def test_stix21_object_mapping_ipv4(self):
        """Test STIX 2.1 ipv4-addr -> our descriptor."""
        import server
        d = server._stix21_object_to_our(
            {'type': 'ipv4-addr', 'value': '203.0.113.9', 'labels': ['c2'], 'confidence': 85})
        self.assertIsNotNone(d)
        self.assertEqual(d['type'], 'ipv4-addr')
        self.assertEqual(d['object']['ipv4-addr']['value'], '203.0.113.9')
        self.assertEqual(d['confidence'], 85)

    def test_stix21_object_mapping_domain(self):
        """Test STIX 2.1 domain-name -> our descriptor."""
        import server
        d = server._stix21_object_to_our({'type': 'domain-name', 'value': 'bad.example.com'})
        self.assertIsNotNone(d)
        self.assertEqual(d['type'], 'domain-name')
        self.assertEqual(d['object']['domain-name']['value'], 'bad.example.com')

    def test_stix21_object_mapping_file(self):
        """Test STIX 2.1 file (hashes) -> our file-hash descriptor."""
        import server
        d = server._stix21_object_to_our(
            {'type': 'file', 'hashes': {'SHA-256': 'b' * 64}})
        self.assertIsNotNone(d)
        self.assertEqual(d['type'], 'file-hash')
        self.assertEqual(d['object']['hash_value']['algorithm'], 'sha256')
        self.assertEqual(d['object']['hash_value']['value'], 'b' * 64)

    def test_stix21_object_mapping_indicator_pattern(self):
        """Test STIX 2.1 indicator (IP pattern) -> our ipv4 descriptor."""
        import server
        d = server._stix21_object_to_our({
            'type': 'indicator', 'pattern': "[ipv4-addr:value = '198.51.100.7']"})
        self.assertIsNotNone(d)
        # Pattern points at an IP -> mapped to ipv4-addr.
        self.assertEqual(d['type'], 'ipv4-addr')
        self.assertEqual(d['object']['ipv4-addr']['value'], '198.51.100.7')

    def test_stix21_object_mapping_unsupported(self):
        """Test unsupported/invalid STIX objects are dropped."""
        import server
        # 'tool' is not a supported store type -> dropped.
        self.assertIsNone(
            server._stix21_object_to_our({'type': 'tool', 'name': 'x'}))
        self.assertIsNone(server._stix21_object_to_our({'type': 'ipv4-addr', 'value': '999.1.1.1'}))
        self.assertIsNone(server._stix21_object_to_our({'type': 'file'}))
        # report without object_refs is invalid -> dropped
        self.assertIsNone(server._stix21_object_to_our(
            {'type': 'report', 'name': 'r'}))

    def test_stix21_object_mapping_community_sdos(self):
        """Community-feed SDOs (report/identity/attack-pattern/vulnerability)."""
        import server
        d = server._stix21_object_to_our({'type': 'attack-pattern', 'name': 'T1059'})
        self.assertEqual(d['type'], 'attack-pattern')
        d = server._stix21_object_to_our({'type': 'vulnerability', 'name': 'CVE-2026-1'})
        self.assertEqual(d['type'], 'vulnerability')
        d = server._stix21_object_to_our({'type': 'identity', 'name': 'Acme'})
        self.assertEqual(d['type'], 'identity')
        d = server._stix21_object_to_our({
            'type': 'report', 'name': 'Daily', 'object_refs': ['indicator--a']})
        self.assertEqual(d['type'], 'report')
        self.assertEqual(d['object']['report']['object_refs'], ['indicator--a'])

    def test_stix21_object_mapping_extended_types(self):
        """Test the extended SCO / SDO / relationship mapping (graph)."""
        import server
        # url
        d = server._stix21_object_to_our({'type': 'url', 'value': 'http://evil.example.com/x'})
        self.assertIsNotNone(d)
        self.assertEqual(d['type'], 'url')
        self.assertEqual(d['object']['url']['value'], 'http://evil.example.com/x')
        # email-addr
        d = server._stix21_object_to_our({'type': 'email-addr', 'value': 'a@evil.example.com'})
        self.assertEqual(d['type'], 'email-addr')
        # invalid email -> dropped
        self.assertIsNone(server._stix21_object_to_our({'type': 'email-addr', 'value': 'not-an-email'}))
        # malware SDO (previously unsupported)
        d = server._stix21_object_to_our({'type': 'malware', 'name': 'trickbot'})
        self.assertEqual(d['type'], 'malware')
        self.assertEqual(d['object']['malware']['name'], 'trickbot')
        # relationship (graph edge preserved)
        d = server._stix21_object_to_our({
            'type': 'relationship',
            'source_ref': 'malware--a',
            'relationship_type': 'uses',
            'target_ref': 'tool--b',
        })
        self.assertEqual(d['type'], 'relationship')
        self.assertEqual(d['object']['relationship']['relationship_type'], 'uses')
        # incomplete relationship -> dropped
        self.assertIsNone(server._stix21_object_to_our({
            'type': 'relationship', 'source_ref': 'malware--a'}))

    def test_taxii_puller_init(self):
        """Test TaxiiPuller initialization and URL building."""
        import server
        p = server.TaxiiPuller({
            'name': 'acme',
            'base_url': 'https://taxii.example.com/taxii2/',
            'username': 'u', 'password': 'p',
            'collection': 'feed1',
            'poll_interval': 120,
            'max_objects_per_poll': 10,
        })
        self.assertEqual(p.name, 'acme')
        self.assertEqual(p.base_url, 'https://taxii.example.com/taxii2')
        self.assertEqual(p.collection, 'feed1')
        self.assertEqual(p.max_objects_per_poll, 10)
        self.assertEqual(p.state_id, 'taxii:acme')
        self.assertFalse(p.enabled)
        # Basic auth header is built.
        self.assertIn('Authorization', p._headers('x'))
        self.assertEqual(p._headers('x')['Authorization'].startswith('Basic '), True)
        # API root url keeps the trailing slash.
        self.assertTrue(p._api_root_url().endswith('/'))

    def test_taxii_puller_parses_bare_bundle(self):
        """The generic puller must parse both the spec envelope and the
        OTX-style bare bundle ({type:"bundle", objects:[…]})."""
        import server
        p = server.TaxiiPuller({'name': 'x', 'base_url': 'https://h/', 'collection': 'c'})
        p._get_taxii = lambda path, since=None: {
            'type': 'bundle', 'more': False,
            'objects': [{'type': 'ipv4-addr', 'value': '1.2.3.4'}]}
        self.assertEqual(len(p._fetch_objects(None)), 1)          # bare bundle
        p._get_taxii = lambda path, since=None: {
            'content': {'content': {'type': 'bundle',
                                    'objects': [{'type': 'ipv4-addr', 'value': '5.6.7.8'}]}}}
        self.assertEqual(len(p._fetch_objects(None)), 1)          # spec envelope

    def test_taxii_puller_misconfigured(self):
        """Test a puller missing base_url/collection reports misconfigured."""
        import server
        p = server.TaxiiPuller({'name': 'x'})
        added, err = p._poll_once()
        self.assertEqual(added, 0)
        self.assertIn('misconfigured', err)

    def test_taxii_puller_fetch_and_merge(self):
        """Test TaxiiPuller fetch + map + merge (with a stubbed _fetch_objects)."""
        import server
        _reset_db()
        p = server.TaxiiPuller({
            'name': 'testpull', 'base_url': 'https://x/taxii2', 'collection': 'c',
        })
        p._fetch_objects = lambda since, limit=None: [
            {'type': 'ipv4-addr', 'value': '8.8.8.8', 'confidence': 60},
            {'type': 'domain-name', 'value': 'c2.example.com'},
            {'type': 'ipv4-addr', 'value': '10.1.1.1'},  # private -> still stored, gated
            {'type': 'tool', 'name': 'drop'},             # unsupported -> skipped
        ]
        added, err = p._poll_once()
        self.assertIsNone(err)
        self.assertEqual(added, 3)
        # They are stored with source='testpull'.
        session = server.create_session()
        try:
            rows = session.query(server.STIXObject).filter_by(source='testpull').all()
            vals = {r.ip_address or r.domain for r in rows}
        finally:
            session.close()
        self.assertEqual(vals, {'8.8.8.8', 'c2.example.com', '10.1.1.1'})
        # State was persisted.
        st = server.read_puller_state(p.state_id)
        self.assertEqual(st['last_status'], 'ok')
        self.assertEqual(st['last_added'], 3)

    def test_community_pullers_endpoint_requires_auth(self):
        """Test /community/pullers requires auth."""
        self.client.delete_cookie('taxii2_ui_session')
        response = self.client.get('/community/pullers')
        self.assertEqual(response.status_code, 401)
        # With TAXII headers it works and lists the OTX puller.
        r2 = self.client.get('/community/pullers', headers=self._get_auth_headers())
        self.assertEqual(r2.status_code, 200)
        names = [p['name'] for p in r2.get_json()['pullers']]
        self.assertIn('AlienVault OTX', names)

    def test_community_pull_unknown(self):
        """Test /community/pull/<name> 404s for an unknown puller."""
        response = self.client.post(
            '/community/pull/nope', headers=self._get_auth_headers())
        self.assertEqual(response.status_code, 404)

    def test_auth_success(self):
        """Test successful authentication"""
        self.client.post(
            '/subscriptions/my_client',
            json={'password': 'test_password'}
        )
        response = self.client.post('/auth', data={
            'username': 'my_client',
            'password': 'test_password',
        })
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertTrue(data['authenticated'])
        self.assertEqual(data['client_id'], 'my_client')

    def test_auth_failure(self):
        """Test failed authentication (wrong password)"""
        self.client.post(
            '/subscriptions/my_client',
            json={'password': 'correct_password'}
        )
        response = self.client.post('/auth', data={
            'username': 'my_client',
            'password': 'wrong_password',
        })
        self.assertEqual(response.status_code, 401)

    def test_list_subscriptions(self):
        """Test list subscriptions endpoint"""
        self.client.post(
            '/subscriptions/ClientA',
            json={'password': 'pass1'}
        )
        response = self.client.get('/subscriptions')
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertIsInstance(data, list)
        self.assertTrue(any(sub['client_id'] == 'ClientA' for sub in data))

    def test_stix_object_to_dicts(self):
        """Test ThreatIntel object conversion to STIX format"""
        # IP address
        ip_obj = ThreatIntel(
            stix_id="ipv4-addr--123",
            object_type="ipv4-addr",
            labels=["threat", "malicious"],
            confidence=80,
            ip_address="192.168.1.100",
        )
        stix = ip_obj.to_stix_object()
        self.assertEqual(stix['id'], "ipv4-addr--123")
        self.assertEqual(stix['type'], "ipv4-addr")
        self.assertEqual(stix['object']['value'], "192.168.1.100")
        self.assertEqual(stix['labels'], ["threat", "malicious"])

        # File hash
        hash_obj = ThreatIntel(
            stix_id="file-hash--abc123",
            object_type="file-hash",
            labels=["sha256"],
            confidence=75,
            hash_value="a" * 64,
        )
        stix_hash = hash_obj.to_stix_object()
        self.assertEqual(stix_hash['hash_value']['algorithm'], "sha256")
        self.assertEqual(stix_hash['hash_value']['value'], "a" * 64)

    def test_generate_taxii_bundle(self):
        """Test TAXII bundle generation"""
        bundles = [
            {
                "id": "ipv4-addr--123",
                "type": "ipv4-addr",
                "object": {
                    "ipv4-addr": {
                        "value": "192.168.1.100",
                        "host_name": None,
                    }
                },
            },
        ]
        xml = generate_taxii_bundle(bundles)
        self.assertIn('Response', xml)
        self.assertIn('192.168.1.100', xml)
        self.assertIn('v2_1', xml)

    def test_otx_poller_init(self):
        """Test OtxPoller initialization"""
        cfg = {
            'base_url': 'https://otx.alienvault.com/',
            'api_key': 'abc123',
            'poll_interval': 60,
            'max_pulses': 10,
            'object_types': ['ipv4-addr', 'domain-name', 'file-hash'],
        }
        poller = OtxPoller(cfg)
        self.assertEqual(poller.base_url, 'https://otx.alienvault.com')
        self.assertEqual(poller.api_key, 'abc123')
        self.assertEqual(poller.max_pulses, 10)
        self.assertEqual(poller.object_types, ['ipv4-addr', 'domain-name', 'file-hash'])
        self.assertFalse(poller.enabled)
        self.assertFalse(poller._stop)
        self.assertEqual(poller.SOURCE, 'otx')

    def test_otx_poller_urls(self):
        """Test OTX URL construction (pulses listing + indicators)"""
        poller = OtxPoller({'base_url': 'https://otx.example.com/',
                            'max_pulses': 42, 'max_indicators_per_pulse': 7})
        captured = {}

        def fake_get_json(url):
            captured[url] = True
            # return an empty results list for both endpoint shapes
            return {'results': []}

        poller._get_json = fake_get_json
        poller.fetch_recent_pulses()
        poller.fetch_pulse_indicators('PID')

        self.assertEqual(len(captured), 2)
        listing = [u for u in captured if '/otxapi/pulses/?' in u]
        indicators = [u for u in captured if '/otxapi/pulses/PID/indicators/' in u]
        self.assertEqual(len(listing), 1)
        self.assertEqual(len(indicators), 1)
        self.assertIn('sort=-modified', listing[0])
        self.assertIn('limit=42', listing[0])
        self.assertIn('limit=7', indicators[0])
        self.assertTrue(listing[0].startswith('https://otx.example.com/otxapi/pulses/'))

    def test_otx_poller_headers(self):
        """Test OTX request headers (API key optional)"""
        with_key = OtxPoller({'api_key': 'k123'})._headers()
        self.assertEqual(with_key['X-OTX-API-KEY'], 'k123')
        self.assertIn('Accept', with_key)

        without_key = OtxPoller({})._headers()
        self.assertNotIn('X-OTX-API-KEY', without_key)

    def test_otx_indicator_mapping_ipv4(self):
        """Test OTX IPv4 indicator -> STIX descriptor"""
        poller = OtxPoller({'object_types': ['ipv4-addr', 'domain-name', 'file-hash']})
        stix = poller.otx_indicator_to_stix(
            {'indicator': '203.0.113.7', 'type': 'IPv4'}, 'TestPulse', 0)
        self.assertIsNotNone(stix)
        self.assertEqual(stix['type'], 'ipv4-addr')
        self.assertEqual(stix['id'], 'ipv4-addr--203-0-113-7')
        self.assertEqual(stix['object']['ipv4-addr']['value'], '203.0.113.7')
        # The pulse name is the provenance label; the `otx` source tag is not
        # duplicated as a label.
        self.assertIn('TestPulse', stix['labels'])
        self.assertNotIn('otx', stix['labels'])
        self.assertEqual(stix['confidence'], 70)

    def test_otx_indicator_mapping_domain(self):
        """Test OTX domain/hostname indicator -> STIX descriptor"""
        poller = OtxPoller({'object_types': ['ipv4-addr', 'domain-name', 'file-hash']})
        stix = poller.otx_indicator_to_stix(
            {'indicator': 'evil.example.com', 'type': 'domain'}, 'P', 0)
        self.assertIsNotNone(stix)
        self.assertEqual(stix['type'], 'domain-name')
        self.assertEqual(stix['id'], 'domain-name--evil-example-com')

        host = poller.otx_indicator_to_stix(
            {'indicator': 'host.example.com', 'type': 'hostname'}, 'P', 0)
        self.assertEqual(host['type'], 'domain-name')

    def test_otx_indicator_mapping_hash(self):
        """Test OTX file hash indicator -> STIX descriptor with algo"""
        poller = OtxPoller({'object_types': ['ipv4-addr', 'domain-name', 'file-hash']})
        stix = poller.otx_indicator_to_stix(
            {'indicator': 'a' * 64, 'type': 'FileHash-SHA256'}, 'P', 0)
        self.assertIsNotNone(stix)
        self.assertEqual(stix['type'], 'file-hash')
        self.assertEqual(stix['object']['hash_value']['algorithm'], 'sha256')
        self.assertEqual(stix['object']['hash_value']['value'], 'a' * 64)
        self.assertIn('sha256', stix['labels'])

    def test_otx_indicator_mapping_skips_unsupported(self):
        """Test OTX indicator types we do not consume are dropped"""
        poller = OtxPoller({'object_types': ['ipv4-addr', 'domain-name', 'file-hash']})
        for bad in (
            {'indicator': 'http://evil.example/x', 'type': 'URL'},
            {'indicator': 'rule { ... }', 'type': 'YARA'},
            {'indicator': '2001:db8::1', 'type': 'IPV6'},
            {'indicator': 'a@b.com', 'type': 'email'},
        ):
            self.assertIsNone(
                poller.otx_indicator_to_stix(bad, 'P', 0), msg=str(bad))

    def test_otx_indicator_mapping_type_filter(self):
        """Test configured object_types filter excludes unwanted types"""
        poller = OtxPoller({'object_types': ['ipv4-addr']})
        stix = poller.otx_indicator_to_stix(
            {'indicator': 'evil.example.com', 'type': 'domain'}, 'P', 0)
        self.assertIsNone(stix)  # domain-name not in object_types

    def test_merge_does_not_wipe_manual(self):
        """Test merge-mode OTX ingest never removes manual intel"""
        _reset_db()
        ingest_objects([_manual_obj('192.168.1.100')], mode='replace', source='manual')
        ingest_objects([_otx_obj('8.8.8.8')], mode='merge', source='otx')

        response = self.client.get('/objects', headers=self._get_auth_headers())
        self.assertEqual(response.status_code, 200)
        values = {o['value'] for o in response.get_json()['objects']}
        self.assertIn('192.168.1.100', values)  # manual survived
        self.assertIn('8.8.8.8', values)        # OTX added

    def test_replace_scoped_to_manual(self):
        """Test manual replace wipes manual only, not OTX rows"""
        _reset_db()
        ingest_objects([_otx_obj('8.8.8.8')], mode='merge', source='otx')
        # Simulate a web-UI save with a DIFFERENT manual entry.
        ingest_objects([_manual_obj('10.0.0.1')], mode='replace', source='manual')

        response = self.client.get('/objects', headers=self._get_auth_headers())
        values = {o['value'] for o in response.get_json()['objects']}
        self.assertIn('10.0.0.1', values)   # new manual entry
        self.assertIn('8.8.8.8', values)    # OTX row survived the UI save

    def test_merge_is_upsert(self):
        """Test merge mode updates an existing row instead of duplicating"""
        _reset_db()
        ingest_objects([_otx_obj('8.8.8.8')], mode='merge', source='otx')
        ingest_objects([_otx_obj('8.8.8.8')], mode='merge', source='otx')

        response = self.client.get('/objects', headers=self._get_auth_headers())
        values = [o['value'] for o in response.get_json()['objects']]
        self.assertEqual(values.count('8.8.8.8'), 1)  # no duplicate
        self.assertEqual(response.get_json()['count'], 1)

    # ------------------------------------------------------------------
    # Intel gate (filter between server and Vision One)
    # ------------------------------------------------------------------
    def _taxii_values(self):
        """Values actually served to Vision One via the TAXII collection."""
        r = self.client.get(
            '/taxii2/collections/threat_intel/objects/',
            headers=self._get_auth_headers(),
        )
        self.assertEqual(r.status_code, 200)
        envelope = r.get_json()
        content = envelope['content']['content']  # STIX 2.1 bundle
        return {served_value(o) for o in content['objects'] if served_value(o)}

    def test_gate_private_ip_withheld(self):
        """Community private IP is withheld from TAXII but still stored."""
        _reset_db()
        # public community IP (served) + private community IP (withheld)
        ingest_objects([_otx_obj('8.8.8.8')], mode='merge', source='otx')
        ingest_objects([
            {'id': 'ipv4-addr--10-0-0-5', 'type': 'ipv4-addr',
             'object': {'ipv4-addr': {'value': '10.0.0.5'}},
             'labels': ['otx'], 'confidence': 70}
        ], mode='merge', source='otx')

        served = self._taxii_values()
        self.assertIn('8.8.8.8', served)      # public community IP served
        self.assertNotIn('10.0.0.5', served)  # private withheld

        # But it is still stored (web UI / audit still see it).
        stored = self.client.get('/objects', headers=self._get_auth_headers()).get_json()
        stored_vals = {o['value'] for o in stored['objects']}
        self.assertIn('10.0.0.5', stored_vals)

    def test_gate_manual_never_filtered(self):
        """Manual private IP is always served (manual is never gated)."""
        _reset_db()
        ingest_objects([
            {'id': 'ipv4-addr--192-168-1-5', 'type': 'ipv4-addr',
             'object': {'ipv4-addr': {'value': '192.168.1.5'}},
             'labels': ['manual'], 'confidence': 90}
        ], mode='replace', source='manual')

        served = self._taxii_values()
        self.assertIn('192.168.1.5', served)  # manual private IP passes

    def test_gate_private_ip_toggle_off(self):
        """When drop_private_ips is off, community private IP is served."""
        import server
        old = server.FILTER_DROP_PRIVATE_IPS
        server.FILTER_DROP_PRIVATE_IPS = False
        try:
            _reset_db()
            ingest_objects([
                {'id': 'ipv4-addr--172-16-5-9', 'type': 'ipv4-addr',
                 'object': {'ipv4-addr': {'value': '172.16.5.9'}},
                 'labels': ['otx'], 'confidence': 70}
            ], mode='merge', source='otx')
            served = self._taxii_values()
            self.assertIn('172.16.5.9', served)
        finally:
            server.FILTER_DROP_PRIVATE_IPS = old

    def test_gate_stale_withheld(self):
        """Community object older than freshness_days is withheld."""
        import server
        old_days = server.FILTER_FRESHNESS_DAYS
        server.FILTER_FRESHNESS_DAYS = 30
        try:
            _reset_db()
            stale = _otx_obj('8.8.4.4')
            ingest_objects([stale], mode='merge', source='otx')
            # Backdate the stored row's last_seen/modified by 40 days.
            session = server.create_session()
            try:
                row = session.query(server.STIXObject).filter_by(
                    stix_id=stale['id']).one()
                row.last_seen = datetime.utcnow() - timedelta(days=40)
                row.modified = datetime.utcnow() - timedelta(days=40)
                session.commit()
            finally:
                session.close()
            # Rehydrate so the memory store matches the backdated row.
            with server.memory_lock:
                for sid in list(server.memory_store):
                    server.memory_store.pop(sid)
            server.rehydrate_memory()

            served = self._taxii_values()
            self.assertNotIn('8.8.4.4', served)  # stale community IP withheld
        finally:
            server.FILTER_FRESHNESS_DAYS = old_days

    def test_gate_freshness_toggle_off(self):
        """When freshness_days is null, a stale community object is served."""
        import server
        old_days = server.FILTER_FRESHNESS_DAYS
        server.FILTER_FRESHNESS_DAYS = None
        try:
            _reset_db()
            stale = _otx_obj('8.8.4.4')
            ingest_objects([stale], mode='merge', source='otx')
            session = server.create_session()
            try:
                row = session.query(server.STIXObject).filter_by(
                    stix_id=stale['id']).one()
                row.last_seen = datetime.utcnow() - timedelta(days=40)
                row.modified = datetime.utcnow() - timedelta(days=40)
                session.commit()
            finally:
                session.close()
            with server.memory_lock:
                for sid in list(server.memory_store):
                    server.memory_store.pop(sid)
            server.rehydrate_memory()

            served = self._taxii_values()
            self.assertIn('8.8.4.4', served)  # freshness off -> served
        finally:
            server.FILTER_FRESHNESS_DAYS = old_days

    def test_gate_confidence_floor(self):
        """Community object below min_confidence is withheld."""
        import server
        old = server._FILTER_MIN_CONFIDENCE
        server._FILTER_MIN_CONFIDENCE = 80
        try:
            _reset_db()
            # OTX community objects default to confidence 70 < 80.
            ingest_objects([_otx_obj('8.8.4.4')], mode='merge', source='otx')
            served = self._taxii_values()
            self.assertNotIn('8.8.4.4', served)
        finally:
            server._FILTER_MIN_CONFIDENCE = old

    def test_gate_blocklist(self):
        """Community value on the blocklist is withheld."""
        import server
        old = server.FILTER_BLOCKLIST
        server.FILTER_BLOCKLIST = {'blocked.example.com'}
        try:
            _reset_db()
            ingest_objects([
                {'id': 'domain-name--blocked-example-com', 'type': 'domain-name',
                 'object': {'domain-name': {'value': 'blocked.example.com'}},
                 'labels': ['otx'], 'confidence': 70},
            ], mode='merge', source='otx')
            served = self._taxii_values()
            self.assertNotIn('blocked.example.com', served)
        finally:
            server.FILTER_BLOCKLIST = old

    def test_gate_reason_unit(self):
        """community_intel_reason returns the right reason (or None)."""
        import server
        from server import ThreatIntel, community_intel_reason

        # private IP -> 'private-ip' (drop_private_ips default True)
        ip_obj = ThreatIntel(stix_id='ipv4-addr--1-2-3-4', object_type='ipv4-addr',
                             ip_address='10.1.2.3', source='otx')
        self.assertEqual(community_intel_reason(ip_obj, _row()), 'private-ip')

        # public IP, fresh, no blocklist -> None (served)
        pub = ThreatIntel(stix_id='ipv4-addr--8-8-8-8', object_type='ipv4-addr',
                          ip_address='8.8.8.8', source='otx', confidence=70)
        self.assertIsNone(community_intel_reason(pub, _row(last_seen=datetime.utcnow())))

    def test_private_ip_helper(self):
        """_is_private_or_reserved_ip classification sanity."""
        import server
        f = server._is_private_or_reserved_ip
        self.assertTrue(f('10.0.0.1'))
        self.assertTrue(f('192.168.1.1'))
        self.assertTrue(f('172.16.0.1'))
        self.assertTrue(f('127.0.0.1'))
        self.assertTrue(f('169.254.1.1'))
        self.assertTrue(f('0.0.0.0'))
        self.assertTrue(f('255.255.255.255'))
        self.assertFalse(f('8.8.8.8'))
        self.assertFalse(f('1.1.1.1'))
        self.assertFalse(f('not-an-ip'))
        self.assertFalse(f(''))

    def test_vision_one_poller_init(self):
        """Test VisionOnePoller initialization"""
        cfg = {
            'base_url': 'http://localhost:5000',
            'client_id': 'vision_one_client',
            'name': 'Trend Micro Vision One',
            'poll_interval': 60,
            'subscription_name': 'Vision One TAXII Feed',
            'object_types': ['ipv4-addr', 'file-hash', 'domain-name'],
        }
        poller = VisionOnePoller(cfg)
        self.assertEqual(poller.base_url, 'http://localhost:5000')
        self.assertEqual(poller.client_id, 'vision_one_client')
        self.assertEqual(poller.subscription_name, 'Vision One TAXII Feed')
        self.assertEqual(poller.object_types, ['ipv4-addr', 'file-hash', 'domain-name'])
        self.assertFalse(poller._stop)

    def test_vision_one_poller_build_url(self):
        """Test Vision One URL construction"""
        cfg = {
            'base_url': 'http://localhost:5000/',
            'client_id': 'vision_one',
            'poll_interval': 60,
            'subscription_name': 'Vision One TAXII Feed',
        }
        poller = VisionOnePoller(cfg)
        url = poller.base_url
        self.assertEqual(url, 'http://localhost:5000')

    def test_vision_one_poller_auth_headers(self):
        """Test Vision One authentication headers"""
        # Configure auth in server
        app.config['TAXII_AUTH'] = {
            'username': 'vis_user',
            'password': 'vis_pass',
        }
        cfg = {
            'base_url': 'http://localhost:5000',
            'client_id': 'vision_one',
        }
        poller = VisionOnePoller(cfg)
        headers = poller._get_auth_headers()
        self.assertEqual(headers['X-Taxii-Username'], 'vis_user')
        self.assertEqual(headers['X-Taxii-Password'], 'vis_pass')


class TestRevocationAndTTL(unittest.TestCase):
    """Manual revocation, TTL auto-revocation, and TAXII 2.1 match filters."""

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        app.config['TAXII_AUTH'] = {
            'username': 'test_taxii_user',
            'password': 'test_taxii_pass',
        }
        # Revoke is a feed-wide write: needs the admin principal in tests.
        app.config['TAXII_ADMIN_AUTH'] = {
            'username': 'test_taxii_user',
            'password': 'test_taxii_pass',
        }
        self.auth = {
            'X-Taxii-Username': 'test_taxii_user',
            'X-Taxii-Password': 'test_taxii_pass',
        }

    def _ingest(self, objs):
        return self.client.post(
            '/feed/ingest', json={'stix_objects': objs}, headers=self.auth
        )

    def _bundle(self, qs=''):
        r = self.client.get(
            '/taxii2/collections/threat_intel/objects/' + qs,
            headers=self.auth,
        )
        self.assertEqual(r.status_code, 200)
        return r.get_json()['content']['content']['objects']

    def test_revoke_object(self):
        self.assertEqual(self._ingest([_manual_obj('1.2.3.4')]).status_code, 200)
        # default action = revoke
        r = self.client.post(
            '/objects/ipv4-addr--1-2-3-4/revoke', json={}, headers=self.auth
        )
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.get_json()['revoked'])
        # served in the TAXII bundle with revoked: true
        objs = self._bundle()
        self.assertEqual(len(objs), 1)
        self.assertTrue(objs[0]['revoked'])

    def test_unrevoke_object(self):
        self._ingest([_manual_obj('5.6.7.8')])
        self.client.post(
            '/objects/ipv4-addr--5-6-7-8/revoke', json={'action': 'revoke'},
            headers=self.auth,
        )
        r = self.client.post(
            '/objects/ipv4-addr--5-6-7-8/revoke', json={'action': 'unrevoke'},
            headers=self.auth,
        )
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.get_json()['revoked'])
        objs = self._bundle()
        self.assertFalse(objs[0]['revoked'])

    def test_revoke_no_auth(self):
        self._ingest([_manual_obj('9.8.7.6')])
        r = self.client.post('/objects/ipv4-addr--9-8-7-6/revoke', json={})
        self.assertEqual(r.status_code, 401)

    def test_revoke_unknown_object(self):
        r = self.client.post(
            '/objects/ipv4-addr--0-0-0-0/revoke', json={}, headers=self.auth
        )
        self.assertEqual(r.status_code, 404)

    def test_revoke_bad_action(self):
        self._ingest([_manual_obj('4.3.2.1')])
        r = self.client.post(
            '/objects/ipv4-addr--4-3-2-1/revoke', json={'action': 'nope'},
            headers=self.auth,
        )
        self.assertEqual(r.status_code, 400)

    def test_revoke_sticky_across_replace_ingest(self):
        """A UI save (replace) must not silently reinstate a revoked object."""
        self._ingest([_manual_obj('2.2.2.2')])
        self.client.post(
            '/objects/ipv4-addr--2-2-2-2/revoke', json={'action': 'revoke'},
            headers=self.auth,
        )
        # save again WITHOUT a revoked flag -> stays revoked
        self._ingest([_manual_obj('2.2.2.2')])
        objs = self._bundle()
        self.assertEqual(len(objs), 1)
        self.assertTrue(objs[0]['revoked'])

    def test_objects_endpoint_reports_revoked(self):
        self._ingest([_manual_obj('3.3.3.3')])
        self.client.post(
            '/objects/ipv4-addr--3-3-3-3/revoke', json={}, headers=self.auth
        )
        r = self.client.get('/objects', headers=self.auth)
        o = r.get_json()['objects'][0]
        self.assertTrue(o['revoked'])
        self.assertIsNotNone(o['revoked_at'])

    def test_ttl_sweep_auto_revokes_expired(self):
        """An IP older than ipv4_days is auto-revoked (served revoked:true)."""
        old = datetime.utcnow() - timedelta(days=30)  # > default 14d
        self._ingest([_manual_obj('6.6.6.6')])
        # backdate the row past its TTL
        session = _server.create_session()
        try:
            row = session.query(STIXObject).filter_by(
                stix_id='ipv4-addr--6-6-6-6').first()
            row.last_seen = old
            row.first_seen = old
            session.commit()
        finally:
            session.close()
        newly = sweep_expired_indicators()
        self.assertEqual(newly, 1)
        objs = self._bundle()
        self.assertEqual(len(objs), 1)
        self.assertTrue(objs[0]['revoked'])

    def test_ttl_sweep_keeps_fresh_objects(self):
        self._ingest([_manual_obj('8.8.8.8')])  # fresh -> untouched
        self.assertEqual(sweep_expired_indicators(), 0)
        objs = self._bundle()
        self.assertFalse(objs[0]['revoked'])

    def test_match_type_filter(self):
        self._ingest([
            _manual_obj('1.1.1.1'),
            {'id': 'domain-name--a-b-c', 'type': 'domain-name',
             'object': {'domain-name': {'value': 'a.b.c'}}, 'labels': [], 'confidence': 50},
        ])
        # Served objects are indicators; match[type] filters by the underlying
        # intel kind (the pattern shows what each one is).
        objs = self._bundle('?match%5Btype%5D=domain-name')  # match[type]=domain-name
        self.assertEqual(len(objs), 1)
        self.assertEqual(objs[0]['type'], 'indicator')
        self.assertEqual(served_kind(objs[0]), 'domain-name')
        self.assertEqual(served_value(objs[0]), 'a.b.c')
        objs = self._bundle('?match%5Btype%5D=ipv4-addr')
        self.assertEqual(len(objs), 1)
        self.assertEqual(served_kind(objs[0]), 'ipv4-addr')
        self.assertEqual(served_value(objs[0]), '1.1.1.1')
        # match[type]=indicator matches every IOC row (all served as indicators)
        objs = self._bundle('?match%5Btype%5D=indicator')
        self.assertEqual(len(objs), 2)

    def test_match_id_filter(self):
        self._ingest([
            _manual_obj('1.1.1.1'),
            _manual_obj('2.2.2.2'),
        ])
        # match[id] filters by the stored id; the served object is its indicator.
        objs = self._bundle('?match%5Bid%5D=ipv4-addr--1-1-1-1')
        self.assertEqual(len(objs), 1)
        self.assertEqual(objs[0]['type'], 'indicator')
        self.assertEqual(served_value(objs[0]), '1.1.1.1')
        # ...and a client can re-fetch by the id it RECEIVED (the served id).
        served_id = objs[0]['id']
        again = self._bundle('?match%5Bid%5D=' + served_id)
        self.assertEqual(len(again), 1)
        self.assertEqual(again[0]['id'], served_id)

    def test_added_after_param(self):
        self._ingest([_manual_obj('7.7.7.7')])
        future = (datetime.utcnow() + timedelta(days=1)).strftime('%Y-%m-%dT%H:%M:%SZ')
        objs = self._bundle('?added_after=' + future)
        self.assertEqual(len(objs), 0)  # nothing modified after "now+1d"


class TestCsvImport(unittest.TestCase):
    """GT-team CSV upload with fuzzy header mapping."""

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        app.config['TAXII_AUTH'] = {
            'username': 'test_taxii_user',
            'password': 'test_taxii_pass',
        }
        # Import is a write: needs the admin principal in tests.
        app.config['TAXII_ADMIN_AUTH'] = {
            'username': 'test_taxii_user',
            'password': 'test_taxii_pass',
        }
        self.auth = {
            'X-Taxii-Username': 'test_taxii_user',
            'X-Taxii-Password': 'test_taxii_pass',
        }

    def _post_csv(self, text):
        return self.client.post(
            '/feed/import-csv', data={'file': (io.BytesIO(text.encode('utf-8')), 'intel.csv')},
            headers=self.auth,
        )

    def test_fuzzy_headers(self):
        """IP_Address / Destination / malicious_domain / sha256 all map.

        Rows carrying both an IP and a domain yield ONE object per value.
        """
        csv_text = (
            "IP_Address,Destination,malicious_domain,sha256,labels,confidence\n"
            "1.2.3.4,,evil.example.com,,c2,90\n"
            ",5.6.7.8,bot.evil.net,,apt,80\n"
            ",,,abcdef0123456789abcdef0123456789abcdef0123456789abcdef0123456789,mw,70\n"
            "999.1.1.1,,,,,99\n"  # invalid IP -> skipped
        )
        r = self._post_csv(csv_text)
        self.assertEqual(r.status_code, 200)
        body = r.get_json()
        self.assertEqual(body['imported'], 5)  # 2+2+1 objects; row 4 skipped
        self.assertEqual(body['skipped_total'], 1)
        # verify stored types
        r = self.client.get('/objects', headers=self.auth)
        by_type = {}
        for o in r.get_json()['objects']:
            by_type.setdefault(o['type'], []).append(o['value'])
        self.assertEqual(by_type['ipv4-addr'], ['1.2.3.4', '5.6.7.8'])
        self.assertEqual(by_type['domain-name'], ['evil.example.com', 'bot.evil.net'])
        self.assertEqual(by_type['file-hash'],
                         ['abcdef0123456789abcdef0123456789abcdef0123456789abcdef0123456789'])

    def test_import_merges_not_replaces(self):
        """CSV import (merge) must not wipe existing manual rows."""
        self.client.post('/feed/ingest', json={'stix_objects': [_manual_obj('9.9.9.9')]},
                         headers=self.auth)
        r = self._post_csv("ip\n1.1.1.1\n")
        self.assertEqual(r.status_code, 200)
        r = self.client.get('/objects', headers=self.auth)
        values = {o['value'] for o in r.get_json()['objects']}
        self.assertEqual(values, {'9.9.9.9', '1.1.1.1'})

    def test_import_no_auth(self):
        r = self.client.post('/feed/import-csv', data={'file': (io.BytesIO(b'ip\n1.1.1.1'), 'a.csv')})
        self.assertEqual(r.status_code, 401)

    def test_import_empty_csv(self):
        r = self._post_csv("")
        self.assertEqual(r.status_code, 400)

    def test_import_no_recognizable_rows(self):
        r = self._post_csv("foo,bar\nhello,world\n")
        self.assertEqual(r.status_code, 400)
        self.assertIn('no recognizable rows', r.get_json()['error'])


class TestSso(unittest.TestCase):
    """Microsoft Entra ID OIDC SSO: PKCE redirect, state/nonce, JWT
    validation (real RS256 with a local key), allow-lists, cookie issue.

    Only the network calls (code exchange) are mocked; token signature
    verification runs for real against a locally generated RSA key.
    """

    _SSO_KEYS = ('SSO_ENABLED', 'SSO_TENANT', 'SSO_CLIENT_ID', 'SSO_CLIENT_SECRET',
                 'SSO_REDIRECT_URI', 'SSO_AUTH_ENDPOINT', 'SSO_ISSUER',
                 'SSO_ALLOW_DOMAINS', 'SSO_ALLOW_UPNS', 'SSO_ALLOW_GROUPS',
                 '_sso_jwks_client')

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        app.config['TAXII_AUTH'] = {
            'username': 'test_taxii_user', 'password': 'test_taxii_pass',
        }
        self._saved = {k: getattr(_server, k) for k in self._SSO_KEYS}
        _server.SSO_ENABLED = True
        _server.SSO_TENANT = 'test-tenant'
        _server.SSO_CLIENT_ID = 'test-client-id'
        _server.SSO_CLIENT_SECRET = 'test-secret'
        _server.SSO_REDIRECT_URI = 'http://localhost:5000/ui/sso/callback'
        _server.SSO_AUTH_ENDPOINT = (
            'https://login.microsoftonline.com/test-tenant/oauth2/v2.0/authorize')
        _server.SSO_ISSUER = ''
        _server.SSO_ALLOW_DOMAINS = ()
        _server.SSO_ALLOW_UPNS = set()
        _server.SSO_ALLOW_GROUPS = set()

        from cryptography.hazmat.primitives.asymmetric import rsa
        priv = rsa.generate_private_key(public_exponent=65537, key_size=2048)
        self._priv = priv
        self._pub = priv.public_key()
        _server._sso_jwks_client = _FakeJwkClient(self._pub)

    def tearDown(self):
        for k, v in self._saved.items():
            setattr(_server, k, v)

    def _id_token(self, **claims):
        base = {
            'iss': 'https://login.microsoftonline.com/test-tenant/v2.0',
            'aud': 'test-client-id',
            'exp': datetime.utcnow() + timedelta(minutes=5),
            'iat': datetime.utcnow(),
            'sub': 'user-sub-123',
            'nonce': 'nonce-abc',
            'preferred_username': 'jdoe@contoso.com',
        }
        base.update(claims)
        return jwt.encode(base, self._priv, algorithm='RS256')

    def _set_oauth_cookie(self, state, nonce, code_verifier):
        token = _server._sso_oauth_state_token(
            {'state': state, 'nonce': nonce, 'code_verifier': code_verifier})
        # The Flask test client drops a manual Cookie header; set_cookie()
        # is the supported way to send cookies.
        self.client.set_cookie('taxii2_ui_sso', token, domain='localhost')

    def test_sso_disabled_404(self):
        _server.SSO_ENABLED = False
        self.assertEqual(self.client.get('/ui/sso').status_code, 404)
        self.assertEqual(self.client.get('/ui/sso/callback').status_code, 404)

    def test_session_reports_sso(self):
        body = self.client.get('/ui/session').get_json()
        self.assertIn('sso', body)
        self.assertTrue(body['sso']['enabled'])

    def test_sso_start_redirects_with_pkce(self):
        r = self.client.get('/ui/sso')
        self.assertEqual(r.status_code, 302)
        loc = r.headers['Location']
        self.assertTrue(loc.startswith(_server.SSO_AUTH_ENDPOINT))
        for frag in ('client_id=test-client-id', 'response_type=code',
                     'code_challenge_method=S256', 'response_mode=query'):
            self.assertIn(frag, loc)
        # oauth state cookie present
        set_cookies = [h for h in r.headers.getlist('Set-Cookie')
                       if h.startswith('taxii2_ui_sso=')]
        self.assertEqual(len(set_cookies), 1)
        raw = set_cookies[0].split(';', 1)[0].split('=', 1)[1]
        oauth = _server._oauth_state_serializer.loads(raw, max_age=600)
        self.assertIn('state', loc)
        # PKCE: challenge must equal S256(verifier)
        import base64, hashlib
        expected = base64.urlsafe_b64encode(
            hashlib.sha256(oauth['code_verifier'].encode()).digest()
        ).rstrip(b'=').decode()
        self.assertIn('code_challenge=' + urllib.parse.quote(expected), loc)

    def test_callback_rejects_state_mismatch(self):
        self._set_oauth_cookie('state-A', 'nonce-abc', 'verifier-1')
        r = self.client.get('/ui/sso/callback?code=abc&state=state-B')
        self.assertEqual(r.status_code, 401)
        self.assertIn('state mismatch', r.get_data(as_text=True))

    def test_callback_rejects_missing_state_cookie(self):
        r = self.client.get('/ui/sso/callback?code=abc&state=state-A')
        self.assertEqual(r.status_code, 401)

    def test_callback_happy_path_sets_session(self):
        import unittest.mock
        tok = self._id_token()
        with unittest.mock.patch.object(
                _server, '_sso_exchange_code',
                return_value={'id_token': tok}) as exch:
            self._set_oauth_cookie('state-X', 'nonce-abc', 'verifier-X')
            r = self.client.get(
                '/ui/sso/callback?code=authcode123&state=state-X',
            )
        self.assertEqual(r.status_code, 302)
        self.assertEqual(r.headers['Location'], '/')
        exch.assert_called_once_with('authcode123', 'verifier-X')
        set_cookies = [h for h in r.headers.getlist('Set-Cookie')]
        sess = [h for h in set_cookies if h.startswith('taxii2_ui_session=')]
        self.assertEqual(len(sess), 1)
        payload = _server._session_serializer.loads(
            sess[0].split(';', 1)[0].split('=', 1)[1])
        self.assertEqual(payload['user'], 'jdoe@contoso.com')

    def test_callback_rejects_bad_signature_audience(self):
        import unittest.mock
        tok = self._id_token(aud='some-other-app')  # wrong audience
        with unittest.mock.patch.object(
                _server, '_sso_exchange_code',
                return_value={'id_token': tok}):
            self._set_oauth_cookie('s1', 'nonce-abc', 'v')
            r = self.client.get(
                '/ui/sso/callback?code=c&state=s1',
            )
        self.assertEqual(r.status_code, 401)
        self.assertIn('token validation failed', r.get_data(as_text=True))

    def test_callback_rejects_nonce_mismatch(self):
        import unittest.mock
        tok = self._id_token(nonce='different-nonce')
        with unittest.mock.patch.object(
                _server, '_sso_exchange_code',
                return_value={'id_token': tok}):
            self._set_oauth_cookie('s2', 'nonce-abc', 'v')
            r = self.client.get(
                '/ui/sso/callback?code=c&state=s2',
            )
        self.assertEqual(r.status_code, 401)
        self.assertIn('token validation failed', r.get_data(as_text=True))

    def test_allowlist_domain(self):
        _server.SSO_ALLOW_DOMAINS = ('contoso.com',)
        ok, _ = _server._sso_user_allowed({'preferred_username': 'jdoe@contoso.com'})
        self.assertTrue(ok)
        ok, reason = _server._sso_user_allowed({'preferred_username': 'eve@evil.com'})
        self.assertFalse(ok)
        self.assertEqual(reason, 'domain not allowed')

    def test_allowlist_upn(self):
        _server.SSO_ALLOW_UPNS = {'jdoe@contoso.com'}
        ok, _ = _server._sso_user_allowed({'preferred_username': 'jdoe@contoso.com'})
        self.assertTrue(ok)
        ok, _ = _server._sso_user_allowed({'preferred_username': 'other@contoso.com'})
        self.assertFalse(ok)

    def test_allowlist_groups(self):
        _server.SSO_ALLOW_GROUPS = {'group-object-1'}
        ok, _ = _server._sso_user_allowed(
            {'preferred_username': 'jdoe@contoso.com',
             'groups': ['group-object-1', 'other']})
        self.assertTrue(ok)
        ok, reason = _server._sso_user_allowed(
            {'preferred_username': 'jdoe@contoso.com', 'groups': ['other']})
        self.assertFalse(ok)
        self.assertEqual(reason, 'not in an allowed Entra group')

    def test_allowlist_callback_denial(self):
        import unittest.mock
        _server.SSO_ALLOW_UPNS = {'onlythis@contoso.com'}
        tok = self._id_token(preferred_username='jdoe@contoso.com')
        with unittest.mock.patch.object(
                _server, '_sso_exchange_code',
                return_value={'id_token': tok}):
            self._set_oauth_cookie('s3', 'nonce-abc', 'v')
            r = self.client.get(
                '/ui/sso/callback?code=c&state=s3',
            )
        self.assertEqual(r.status_code, 401)
        self.assertIn('not authorized', r.get_data(as_text=True))


class TestDotenv(unittest.TestCase):
    """Minimal .env loader: parsing, precedence, missing file, path override."""

    def setUp(self):
        import tempfile
        self.tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.tmp.name, 'test.env')
        # Keys the tests touch, saved/restored so nothing leaks between tests.
        self._keys = ('TEST_DOTENV_A', 'TEST_DOTENV_B', 'TEST_DOTENV_C',
                      'TAXII_ENV_FILE', 'TEST_DOTENV_OVERRIDE')
        self._saved = {k: os.environ.pop(k, None) for k in self._keys}

    def tearDown(self):
        for k, v in self._saved.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        self.tmp.cleanup()

    def _write(self, text):
        with open(self.path, 'w', encoding='utf-8') as fh:
            fh.write(text)

    def _load(self):
        _server._load_dotenv(_server.Path(self.path))

    def test_parses_values_comments_and_quotes(self):
        self._write(
            '# a comment line\n'
            '\n'
            'TEST_DOTENV_A=plain\n'
            'TEST_DOTENV_B="double quoted"\n'
            "TEST_DOTENV_C='single quoted'\n"
            '   # indented comment\n'
        )
        self._load()
        self.assertEqual(os.environ['TEST_DOTENV_A'], 'plain')
        self.assertEqual(os.environ['TEST_DOTENV_B'], 'double quoted')
        self.assertEqual(os.environ['TEST_DOTENV_C'], 'single quoted')

    def test_process_env_always_wins(self):
        self._write('TEST_DOTENV_A=from-file\n')
        os.environ['TEST_DOTENV_A'] = 'from-env'
        self._load()
        self.assertEqual(os.environ['TEST_DOTENV_A'], 'from-env')

    def test_missing_file_is_noop(self):
        # Should not raise, and should not set anything.
        _server._load_dotenv(_server.Path(self.tmp.name) / 'does-not-exist.env')
        self.assertNotIn('TEST_DOTENV_A', os.environ)

    def test_dotenv_path_default_and_override(self):
        # Default: .env next to server.py (repo root), unless TAXII_ENV_FILE.
        self._write('TEST_DOTENV_A=x\n')
        os.environ['TAXII_ENV_FILE'] = self.path
        try:
            self.assertEqual(str(_server._dotenv_path()), self.path)
        finally:
            del os.environ['TAXII_ENV_FILE']
        default = str(_server._dotenv_path())
        # Default .env lives next to server.py, whatever the repo dir is named.
        self.assertEqual(
            default,
            os.path.join(os.path.dirname(os.path.abspath(_server.__file__)), '.env'),
        )
        self.assertFalse(default.startswith(self.tmp.name))


class TestSaveAndDelete(unittest.TestCase):
    """Web-UI save semantics (mode=merge) + POST /feed/delete (purge option)."""

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        app.config['TAXII_AUTH'] = {
            'username': 'test_taxii_user', 'password': 'test_taxii_pass',
        }
        # /feed/delete is a feed-wide write: needs the admin principal.
        app.config['TAXII_ADMIN_AUTH'] = {
            'username': 'test_taxii_user', 'password': 'test_taxii_pass',
        }
        self.auth = {
            'X-Taxii-Username': 'test_taxii_user',
            'X-Taxii-Password': 'test_taxii_pass',
        }

    def _seed(self):
        # 10 manual IPs + 1 OTX IP (community source, must survive saves).
        manual = [_manual_obj(f"203.0.113.{i}") for i in range(1, 11)]
        self._ingest(manual)
        # Simulate the puller's write path (tags source='otx', merge mode).
        ingest_objects([_otx_obj("8.8.8.8")], mode='merge', source='otx')

    def _ingest(self, stix_objects, mode=None):
        payload = {'stix_objects': stix_objects}
        if mode:
            payload['mode'] = mode
        return self.client.post('/feed/ingest', json=payload, headers=self.auth)

    def _ids(self):
        r = self.client.get('/objects', headers=self.auth)
        self.assertEqual(r.status_code, 200)
        return {o['id']: o for o in r.get_json()['objects']}

    def test_ingest_mode_merge_preserves_everything(self):
        self._seed()
        ids = self._ids()
        self.assertEqual(len(ids), 11)
        self.assertEqual(ids['ipv4-addr--8-8-8-8']['source'], 'otx')

        # UI save: merge with 5 (edited) manual rows + 1 brand-new row.
        save = [_manual_obj(f"203.0.113.{i}") for i in range(1, 6)]
        save.append(_manual_obj("203.0.113.99"))  # new
        r = self._ingest(save, mode='merge')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['mode'], 'merge')

        ids = self._ids()
        # The 5 other manual rows + the OTX row survived (not wiped).
        for i in range(6, 11):
            self.assertIn(f"ipv4-addr--203-0-113-{i}", ids)
        self.assertIn('ipv4-addr--8-8-8-8', ids)
        self.assertEqual(ids['ipv4-addr--8-8-8-8']['source'], 'otx')
        # New row merged in; total = 10 manual + 1 new + 1 otx.
        self.assertIn('ipv4-addr--203-0-113-99', ids)
        self.assertEqual(len(ids), 12)

    def test_ingest_default_is_replace(self):
        self._seed()
        self._ingest([_manual_obj("203.0.113.1")])  # legacy replace
        ids = self._ids()
        # Other manual rows wiped; otx row survives.
        self.assertNotIn('ipv4-addr--203-0-113-2', ids)
        self.assertIn('ipv4-addr--203-0-113-1', ids)
        self.assertIn('ipv4-addr--8-8-8-8', ids)

    def test_ingest_bad_mode_400(self):
        r = self._ingest([_manual_obj()], mode='bogus')
        self.assertEqual(r.status_code, 400)

    def test_delete_ids(self):
        self._seed()
        r = self.client.post('/feed/delete',
                             json={'ids': ['ipv4-addr--203-0-113-1', 'ipv4-addr--8-8-8-8']},
                             headers=self.auth)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['deleted'], 2)
        ids = self._ids()
        self.assertNotIn('ipv4-addr--203-0-113-1', ids)
        self.assertNotIn('ipv4-addr--8-8-8-8', ids)  # any-source delete
        self.assertIn('ipv4-addr--203-0-113-2', ids)

    def test_delete_empty_ids_400(self):
        r = self.client.post('/feed/delete', json={'ids': []}, headers=self.auth)
        self.assertEqual(r.status_code, 400)
        r = self.client.post('/feed/delete', json={}, headers=self.auth)
        self.assertEqual(r.status_code, 400)

    def test_delete_unauth_401(self):
        r = self.client.post('/feed/delete', json={'ids': ['x']})
        self.assertEqual(r.status_code, 401)


class TestExtendedTypesAndValidation(unittest.TestCase):
    """New STIX object types (url / email-addr / ipv6-addr / mac-addr /
    windows-registry-key / autonomous-system), the STIX graph objects
    (relationship / malware / threat-actor / campaign) and ingest validation.
    """

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        app.config['TAXII_AUTH'] = {
            'username': 'test_taxii_user',
            'password': 'test_taxii_pass',
        }
        # Ingest is a write: needs the admin principal in tests.
        app.config['TAXII_ADMIN_AUTH'] = {
            'username': 'test_taxii_user',
            'password': 'test_taxii_pass',
        }
        self.auth = {
            'X-Taxii-Username': 'test_taxii_user',
            'X-Taxii-Password': 'test_taxii_pass',
        }

    def _ingest(self, objs):
        return self.client.post(
            '/feed/ingest', json={'stix_objects': objs}, headers=self.auth)

    def _bundle(self, qs=''):
        r = self.client.get(
            '/taxii2/collections/threat_intel/objects/' + qs,
            headers=self.auth,
        )
        self.assertEqual(r.status_code, 200)
        return r.get_json()['content']['content']['objects']

    def test_ingest_and_serve_url_email_ipv6(self):
        r = self._ingest([
            {'id': 'url--evil-x', 'type': 'url',
             'object': {'url': {'value': 'http://evil.example.com/x'}},
             'labels': ['malicious'], 'confidence': 80},
            {'id': 'email-addr--a-b', 'type': 'email-addr',
             'object': {'email-addr': {'value': 'a@evil.example.com'}},
             'labels': ['malicious'], 'confidence': 60},
            {'id': 'ipv6-addr--2001-0db8--1', 'type': 'ipv6-addr',
             'object': {'ipv6-addr': {'value': '2001:db8::1'}},
             'labels': ['malicious'], 'confidence': 70},
        ])
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['objects_count'], 3)

        objs = self._bundle()
        by_kind = {served_kind(o): o for o in objs}
        self.assertEqual(served_value(by_kind['url']), 'http://evil.example.com/x')
        self.assertEqual(served_value(by_kind['email-addr']), 'a@evil.example.com')
        self.assertEqual(served_value(by_kind['ipv6-addr']), '2001:db8::1')
        self.assertTrue(all(o['type'] == 'indicator' for o in objs))

        # /objects lists them with their value
        r = self.client.get('/objects', headers=self.auth)
        by_type = {o['type']: o for o in r.get_json()['objects']}
        self.assertEqual(by_type['url']['value'], 'http://evil.example.com/x')

    def test_windows_registry_key_and_asn(self):
        r = self._ingest([
            {'id': 'windows-registry-key--hkcu', 'type': 'windows-registry-key',
             'object': {'windows-registry-key': {'key': 'HKEY_CURRENT_USER\\Software\\Evil'}},
             'labels': ['registry'], 'confidence': 55},
            {'id': 'autonomous-system--12345', 'type': 'autonomous-system',
             'object': {'autonomous-system': {'number': 64512}},
             'labels': ['asn'], 'confidence': 50},
        ])
        self.assertEqual(r.status_code, 200)
        objs = self._bundle()
        by_kind = {served_kind(o): o for o in objs}
        self.assertEqual(by_kind['windows-registry-key']['pattern'],
                         "[windows-registry-key:key = 'HKEY_CURRENT_USER\\Software\\Evil']")
        self.assertEqual(by_kind['autonomous-system']['pattern'],
                         "[autonomous-system:number = '64512']")

    def test_indicator_over_url_renders_url_pattern(self):
        self._ingest([
            {'id': 'indicator--cafe', 'type': 'indicator',
             'object': {'indicator': {'value': 'http://evil.example.com/x'}},
             'labels': ['url'], 'confidence': 90},
        ])
        objs = self._bundle()
        self.assertEqual(len(objs), 1)
        o = objs[0]
        self.assertEqual(o['type'], 'indicator')
        self.assertEqual(o['pattern'], "[url:value = 'http://evil.example.com/x']")

    def test_guess_object_type_extended(self):
        import server
        cases = {
            'http://a.example.com/x': ('url', 'url'),
            'https://a.example.com/x?q=1': ('url', 'url'),
            'a@evil.example.com': ('email-addr', 'email-addr'),
            '2001:db8::1': ('ipv6-addr', 'ipv6-addr'),
            'aa:bb:cc:dd:ee:ff': ('mac-addr', 'mac-addr'),
            'AS12345': ('autonomous-system', 'autonomous-system'),
            '1.2.3.4': ('ipv4-addr', 'ipv4-addr'),
            'evil.example.com': ('domain-name', 'domain-name'),
            'free text': ('indicator', 'indicator'),
        }
        for value, (otype, algo_marker) in cases.items():
            guessed, _ = server._guess_object_type(value)
            self.assertEqual(guessed, otype, f'value {value!r}')

    def test_validation_rejects_bad_id_prefix(self):
        # type url but id says ipv4-addr -> skipped, not stored
        r = self._ingest([
            {'id': 'ipv4-addr--oops', 'type': 'url',
             'object': {'url': {'value': 'http://evil.example.com/x'}},
             'labels': ['x'], 'confidence': 50},
        ])
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['objects_count'], 0)
        self.assertEqual(self.client.get('/objects', headers=self.auth).get_json()['count'], 0)

    def test_validation_rejects_unsupported_type_and_bad_confidence(self):
        r = self._ingest([
            {'id': 'banana--1', 'type': 'banana',
             'object': {'banana': {'value': 'x'}}, 'labels': [], 'confidence': 50},
            {'id': 'ipv4-addr--1-2-3-4', 'type': 'ipv4-addr',
             'object': {'ipv4-addr': {'value': '1.2.3.4'}},
             'labels': ['x'], 'confidence': 999},
            {'id': 'ipv4-addr--5-6-7-8', 'type': 'ipv4-addr',
             'object': {'ipv4-addr': {'value': '5.6.7.8'}},
             'labels': ['x'], 'confidence': 50, 'modified': 'not-a-date'},
        ])
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['objects_count'], 0)

    def test_validation_skips_only_bad_rows(self):
        r = self._ingest([
            {'id': 'banana--1', 'type': 'banana',
             'object': {'banana': {'value': 'x'}}, 'labels': [], 'confidence': 50},
            {'id': 'ipv4-addr--9-9-9-9', 'type': 'ipv4-addr',
             'object': {'ipv4-addr': {'value': '9.9.9.9'}},
             'labels': ['x'], 'confidence': 50},
        ])
        self.assertEqual(r.get_json()['objects_count'], 1)
        objs = self.client.get('/objects', headers=self.auth).get_json()['objects']
        self.assertEqual([o['value'] for o in objs], ['9.9.9.9'])

    def test_graph_relationship_ingest_and_serve(self):
        r = self._ingest([
            {'id': 'malware--trickbot', 'type': 'malware',
             'object': {'malware': {'name': 'TrickBot'}},
             'labels': ['malware'], 'confidence': 70},
            {'id': 'threat-actor--apt41', 'type': 'threat-actor',
             'object': {'threat-actor': {'name': 'APT41'}},
             'labels': ['apt'], 'confidence': 70},
            {'id': 'relationship--r1', 'type': 'relationship',
             'object': {'relationship': {
                 'relationship_type': 'uses',
                 'source_ref': 'malware--trickbot',
                 'target_ref': 'threat-actor--apt41'}},
             'labels': ['graph'], 'confidence': 70},
        ])
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['objects_count'], 3)

        objs = self._bundle()
        by_type = {o['type']: o for o in objs}
        self.assertEqual(by_type['malware']['name'], 'TrickBot')
        self.assertTrue(by_type['malware']['is_family'])
        self.assertEqual(by_type['threat-actor']['name'], 'APT41')
        self.assertEqual(by_type['threat-actor']['threat_actor_types'], ['unknown'])
        rel = by_type['relationship']
        self.assertEqual(rel['relationship_type'], 'uses')
        self.assertEqual(rel['source_ref'], 'malware--trickbot')
        self.assertEqual(rel['target_ref'], 'threat-actor--apt41')

        # /objects shows the edge as a readable summary
        r = self.client.get('/objects', headers=self.auth)
        rel_row = next(o for o in r.get_json()['objects'] if o['type'] == 'relationship')
        self.assertIn('malware--trickbot → uses → threat-actor--apt41', rel_row['value'])

    def test_incomplete_relationship_rejected(self):
        r = self._ingest([
            {'id': 'relationship--r1', 'type': 'relationship',
             'object': {'relationship': {'source_ref': 'malware--a'}},
             'labels': [], 'confidence': 50},
        ])
        self.assertEqual(r.get_json()['objects_count'], 0)

    def test_campaign_ingest(self):
        r = self._ingest([
            {'id': 'campaign--c1', 'type': 'campaign',
             'object': {'campaign': {'name': 'Operation Beep'}},
             'labels': ['campaign'], 'confidence': 60},
        ])
        objs = self._bundle()
        self.assertEqual(objs[0]['type'], 'campaign')
        self.assertEqual(objs[0]['name'], 'Operation Beep')


class TestCommunitySdoAndFidelity(unittest.TestCase):
    """Community-feed SDOs (report/identity/attack-pattern/vulnerability) and
    upstream provenance (valid_from) preserved through ingest -> serve."""

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        app.config['TAXII_AUTH'] = {'username': 'u', 'password': 'p'}
        # Ingest is a write: needs the admin principal in tests.
        app.config['TAXII_ADMIN_AUTH'] = {'username': 'u', 'password': 'p'}
        self.auth = {'X-Taxii-Username': 'u', 'X-Taxii-Password': 'p'}

    def _served(self):
        r = self.client.get('/taxii2/collections/threat_intel/objects/',
                            headers=self.auth)
        self.assertEqual(r.status_code, 200)
        return r.get_json()['content']['content']['objects']

    def test_serve_community_sdos(self):
        r = self.client.post('/feed/ingest', json={'stix_objects': [
            {'id': 'attack-pattern--T1059', 'type': 'attack-pattern',
             'object': {'attack-pattern': {'name': 'Command and Scripting Interpreter'}},
             'labels': ['attack'], 'confidence': 60},
            {'id': 'vulnerability--CVE-2026-1', 'type': 'vulnerability',
             'object': {'vulnerability': {'name': 'CVE-2026-1'}},
             'labels': ['cve'], 'confidence': 60},
            {'id': 'identity--acme', 'type': 'identity',
             'object': {'identity': {'name': 'Acme Corp'}},
             'labels': ['identity'], 'confidence': 60},
            {'id': 'report--daily', 'type': 'report',
             'object': {'report': {'name': 'Daily report',
                                   'object_refs': ['indicator--x']}},
             'labels': ['report'], 'confidence': 60},
        ]}, headers=self.auth)
        self.assertEqual(r.get_json()['objects_count'], 4)
        by = {o['type']: o for o in self._served()}
        self.assertEqual(by['attack-pattern']['name'], 'Command and Scripting Interpreter')
        self.assertEqual(by['vulnerability']['name'], 'CVE-2026-1')
        self.assertEqual(by['identity']['identity_class'], 'unknown')   # REQUIRED field
        self.assertEqual(by['report']['object_refs'], ['indicator--x'])  # REQUIRED field

    def test_upstream_valid_from_preserved(self):
        self.client.post('/feed/ingest', json={'stix_objects': [
            {'id': 'indicator--vf', 'type': 'indicator',
             'object': {'indicator': {'value': 'evil.example.com'}},
             'labels': ['x'], 'confidence': 70,
             'valid_from': '2026-09-01T01:00:36.000Z'},
        ]}, headers=self.auth)
        objs = self._served()
        self.assertEqual(objs[0]['type'], 'indicator')
        self.assertEqual(objs[0]['valid_from'], '2026-09-01T01:00:36Z')

    def test_report_without_refs_skipped(self):
        r = self.client.post('/feed/ingest', json={'stix_objects': [
            {'id': 'report--x', 'type': 'report',
             'object': {'report': {'name': 'no refs'}}, 'labels': [], 'confidence': 50},
        ]}, headers=self.auth)
        self.assertEqual(r.get_json()['objects_count'], 0)

    def test_sdo_ids_are_valid_stix(self):
        """A served SDO/SRO id must be <type>--<uuid> even if ingested with a
        human id (strict STIX consumers reject anything else)."""
        import uuid as _uuid
        self.client.post('/feed/ingest', json={'stix_objects': [
            {'id': 'attack-pattern--human', 'type': 'attack-pattern',
             'object': {'attack-pattern': {'name': 'T1059'}},
             'labels': [], 'confidence': 50},
            {'id': 'malware--trickbot', 'type': 'malware',
             'object': {'malware': {'name': 'TrickBot'}},
             'labels': [], 'confidence': 50},
        ]}, headers=self.auth)
        for o in self._served():
            self.assertTrue(o['id'].startswith(o['type'] + '--'), o['id'])
            _uuid.UUID(o['id'].split('--', 1)[1])   # raises if not a UUID


class TestTaxiiPagination(unittest.TestCase):
    """TAXII 2.1 Get Objects pagination: limit / next / more (keyset)."""

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        app.config['TAXII_AUTH'] = {
            'username': 'test_taxii_user',
            'password': 'test_taxii_pass',
        }
        # Ingest is a write: needs the admin principal in tests.
        app.config['TAXII_ADMIN_AUTH'] = {
            'username': 'test_taxii_user',
            'password': 'test_taxii_pass',
        }
        self.auth = {
            'X-Taxii-Username': 'test_taxii_user',
            'X-Taxii-Password': 'test_taxii_pass',
        }

    def _ingest_n(self, n):
        objs = [{
            'id': f'ipv4-addr--1-1-1-{i}',
            'type': 'ipv4-addr',
            'object': {'ipv4-addr': {'value': f'1.1.1.{i}'}},
            'labels': ['t'], 'confidence': 50,
        } for i in range(1, n + 1)]
        r = self.client.post('/feed/ingest', json={'stix_objects': objs},
                             headers=self.auth)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['objects_count'], n)

    def _page(self, qs):
        r = self.client.get(
            '/taxii2/collections/threat_intel/objects/' + qs,
            headers=self.auth)
        self.assertEqual(r.status_code, 200)
        return r.get_json()

    def test_pages_walk_without_duplicates_or_loss(self):
        self._ingest_n(7)
        vals = []
        more = True
        nxt = None
        pages = 0
        while more:
            qs = '?limit=3'
            if nxt:
                qs += '&next=' + urllib.parse.quote(nxt)
            msg = self._page(qs)
            objs = msg['content']['content']['objects']
            self.assertLessEqual(len(objs), 3)
            vals.extend(served_value(o) for o in objs)
            more = msg['more']
            nxt = msg.get('next')
            pages += 1
            self.assertLess(pages, 10)  # guard against infinite loops
        self.assertEqual(pages, 3)
        self.assertEqual(sorted(vals), sorted(f'1.1.1.{i}' for i in range(1, 8)))
        self.assertEqual(len(set(vals)), 7)

    def test_single_page_when_fewer_than_limit(self):
        self._ingest_n(2)
        msg = self._page('?limit=5')
        self.assertFalse(msg['more'])
        self.assertNotIn('next', msg)
        self.assertEqual(len(msg['content']['content']['objects']), 2)

    def test_more_flag_off_without_limit(self):
        self._ingest_n(4)
        msg = self._page('')
        self.assertFalse(msg['more'])
        self.assertNotIn('next', msg)
        self.assertEqual(len(msg['content']['content']['objects']), 4)

    def test_no_limit_still_capped_and_paginates(self):
        """A poll with no ?limit= is bounded by taxii.max_objects.

        Without the cap a single request could materialise a whole collection.
        The client is expected to follow more/next.
        """
        saved = _server.TAXII_CFG.get('max_objects')
        _server.TAXII_CFG['max_objects'] = 3
        try:
            self._ingest_n(7)
            msg = self._page('')
            self.assertEqual(len(msg['content']['content']['objects']), 3)
            self.assertTrue(msg['more'])
            self.assertTrue(msg.get('next'))
            vals = [served_value(o)
                    for o in msg['content']['content']['objects']]
            nxt, guard = msg.get('next'), 0
            while nxt and guard < 10:
                msg = self._page('?next=' + urllib.parse.quote(nxt))
                vals.extend(served_value(o)
                            for o in msg['content']['content']['objects'])
                nxt = msg.get('next') if msg['more'] else None
                guard += 1
            self.assertEqual(sorted(vals),
                             sorted(f'1.1.1.{i}' for i in range(1, 8)))
        finally:
            if saved is None:
                _server.TAXII_CFG.pop('max_objects', None)
            else:
                _server.TAXII_CFG['max_objects'] = saved

    def test_limit_zero_and_non_numeric_400(self):
        self._ingest_n(3)
        for bad in ('0', '-1', 'abc', '1.5'):
            r = self.client.get(
                '/taxii2/collections/threat_intel/objects/?limit=' + bad,
                headers=self.auth)
            self.assertEqual(r.status_code, 400, f'limit={bad}')

    def test_bad_since_and_bad_next_400(self):
        self._ingest_n(2)
        r = self.client.get(
            '/taxii2/collections/threat_intel/objects/?since=garbage',
            headers=self.auth)
        self.assertEqual(r.status_code, 400)
        r = self.client.get(
            '/taxii2/collections/threat_intel/objects/?next=%%%',
            headers=self.auth)
        self.assertEqual(r.status_code, 400)

    def test_pagination_respects_match_type(self):
        self._ingest_n(3)
        self.client.post('/feed/ingest', json={'stix_objects': [
            {'id': 'domain-name--a-b', 'type': 'domain-name',
             'object': {'domain-name': {'value': 'a.b'}},
             'labels': [], 'confidence': 50},
            {'id': 'domain-name--c-d', 'type': 'domain-name',
             'object': {'domain-name': {'value': 'c.d'}},
             'labels': [], 'confidence': 50},
        ]}, headers=self.auth)
        msg = self._page('?limit=1&match%5Btype%5D=domain-name')
        objs = msg['content']['content']['objects']
        self.assertEqual(len(objs), 1)
        self.assertTrue(msg['more'])
        # follow to the end — only domains ever returned
        more, nxt = msg['more'], msg.get('next')
        seen = [served_value(o) for o in objs]
        while more:
            msg = self._page('?limit=1&match%5Btype%5D=domain-name&next='
                             + urllib.parse.quote(nxt))
            seen.extend(served_value(o)
                        for o in msg['content']['content']['objects'])
            more, nxt = msg['more'], msg.get('next')
        self.assertEqual(sorted(seen), ['a.b', 'c.d'])

    def test_pagination_not_changed_by_since_for_old_rows(self):
        # rows ingested at the same batch share a timestamp; the cursor must
        # still separate them by stix_id (no skipped/duplicate rows).
        self._ingest_n(4)
        msg = self._page('?limit=2')
        objs1 = msg['content']['content']['objects']
        msg2 = self._page('?limit=2&next=' + urllib.parse.quote(msg['next']))
        objs2 = msg2['content']['content']['objects']
        all_ids = [o['id'] for o in objs1] + [o['id'] for o in objs2]
        self.assertEqual(len(set(all_ids)), 4)
        self.assertFalse(msg2['more'])

    def test_multi_value_match_filters(self):
        """match[type]/match[id] accept comma-separated lists (spec 5.3):
        every listed value is a disjunction (a real client joins them this
        way — e.g. taxii2client sends match[type]=domain-name,file-hash)."""
        self._ingest_n(3)
        self.client.post('/feed/ingest', json={'stix_objects': [
            {'id': 'domain-name--a-b', 'type': 'domain-name',
             'object': {'domain-name': {'value': 'a.b'}},
             'labels': [], 'confidence': 50},
        ], 'mode': 'merge'}, headers=self.auth)
        # comma-separated match[type] -> union of the listed types
        objs = self._page(
            '?match%5Btype%5D=domain-name,ipv4-addr')['content']['content']['objects']
        self.assertEqual(len(objs), 4)  # 3 ipv4 + 1 domain
        # comma-separated match[id] -> union of the listed ids
        objs = self._page(
            '?match%5Bid%5D=ipv4-addr--1-1-1-1,domain-name--a-b'
        )['content']['content']['objects']
        self.assertEqual(sorted(served_value(o) for o in objs), ['1.1.1.1', 'a.b'])
        # a non-matching single value is still a real filter (not a no-op)
        objs = self._page('?match%5Btype%5D=nope')['content']['content']['objects']
        self.assertEqual(len(objs), 0)


class TestXlsxImport(unittest.TestCase):
    """Excel (.xlsx) import through /feed/import-csv + column typing."""

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        app.config['TAXII_AUTH'] = {
            'username': 'test_taxii_user',
            'password': 'test_taxii_pass',
        }
        # Import is a write: needs the admin principal in tests.
        app.config['TAXII_ADMIN_AUTH'] = {
            'username': 'test_taxii_user',
            'password': 'test_taxii_pass',
        }
        self.auth = {
            'X-Taxii-Username': 'test_taxii_user',
            'X-Taxii-Password': 'test_taxii_pass',
        }

    def _xlsx_bytes(self, rows):
        """Build an .xlsx workbook in memory: rows[0] = headers."""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        for row in rows:
            ws.append([c if c is not None else '' for c in row])
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        return buf.getvalue()

    def _post_xlsx(self, rows, filename='intel.xlsx'):
        return self.client.post(
            '/feed/import-csv',
            data={'file': (io.BytesIO(self._xlsx_bytes(rows)), filename)},
            headers=self.auth,
        )

    def test_xlsx_upload_fuzzy_headers(self):
        # Note: "Destination" maps to the IP role (dstip) — a domain there is
        # dropped, exactly like the CSV path. Domains need a 'domain' header.
        r = self._post_xlsx([
            ['IP_Address', 'domain', 'labels', 'confidence'],
            ['1.2.3.4', 'evil.example.com', 'c2', 90],
            ['5.6.7.8', 'bot.evil.net', 'apt', 80],
        ])
        self.assertEqual(r.status_code, 200, r.get_json())
        body = r.get_json()
        self.assertEqual(body['imported'], 4)
        self.assertEqual(body['skipped_total'], 0)
        r = self.client.get('/objects', headers=self.auth)
        objs = r.get_json()['objects']
        self.assertEqual({o['value'] for o in objs},
                         {'1.2.3.4', '5.6.7.8', 'evil.example.com', 'bot.evil.net'})

    def test_xlsx_indicator_column_promotes_url(self):
        r = self._post_xlsx([
            ['indicator', 'labels', 'confidence'],
            ['http://evil.example.com/x', 'url-ioc', 70],
            ['free text note', 'note', 50],
        ])
        self.assertEqual(r.status_code, 200)
        objs = self.client.get('/objects', headers=self.auth).get_json()['objects']
        by_type = {o['type']: o for o in objs}
        self.assertEqual(by_type['url']['value'], 'http://evil.example.com/x')
        self.assertEqual(by_type['indicator']['value'], 'free text note')

    def test_csv_indicator_column_promotes_url_and_ipv6(self):
        r = self.client.post(
            '/feed/import-csv',
            data={'file': (io.BytesIO(
                ('indicator,labels\n'
                 'http://evil.example.com/x,c2\n'
                 '2001:db8::1,ip\n').encode()), 'intel.csv')},
            headers=self.auth,
        )
        self.assertEqual(r.status_code, 200)
        objs = self.client.get('/objects', headers=self.auth).get_json()['objects']
        types = {o['type'] for o in objs}
        self.assertIn('url', types)
        self.assertIn('ipv6-addr', types)

    def test_xlsx_ip_column_accepts_ipv6(self):
        r = self._post_xlsx([
            ['ip'],
            ['2001:db8::42'],
        ])
        self.assertEqual(r.status_code, 200)
        objs = self.client.get('/objects', headers=self.auth).get_json()['objects']
        self.assertEqual(len(objs), 1)
        self.assertEqual(objs[0]['type'], 'ipv6-addr')

    def test_xlsx_bad_file_400(self):
        r = self.client.post(
            '/feed/import-csv',
            data={'file': (io.BytesIO(b'this is not a real xlsx'), 'intel.xlsx')},
            headers=self.auth,
        )
        # unparseable workbook -> no objects -> 400 (or openpyxl error)
        self.assertEqual(r.status_code, 400)


class TestUiStats(unittest.TestCase):
    """GET /ui/stats — aggregates behind the Intel graph dashboard."""

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        app.config['TAXII_AUTH'] = {'username': 'u', 'password': 'p'}
        self.auth = {'X-Taxii-Username': 'u', 'X-Taxii-Password': 'p'}

    def _seed(self):
        ingest_objects([
            {'id': 'ipv4-addr--1-1-1-1', 'type': 'ipv4-addr',
             'object': {'ipv4-addr': {'value': '1.1.1.1'}},
             'labels': ['otx', 'pulseA'], 'confidence': 70},
            {'id': 'domain-name--a-b', 'type': 'domain-name',
             'object': {'domain-name': {'value': 'a.b'}},
             'labels': ['otx'], 'confidence': 80},
        ], mode='merge', source='otx')
        ingest_objects([
            {'id': 'ipv4-addr--8-8-8-8', 'type': 'ipv4-addr',
             'object': {'ipv4-addr': {'value': '8.8.8.8'}},
             'labels': ['manual'], 'confidence': 90},
        ], mode='merge', source='manual')

    def test_scopes_and_aggregates(self):
        self._seed()
        d = self.client.get('/ui/stats?scope=community', headers=self.auth).get_json()
        self.assertEqual(d['scope'], 'community')
        self.assertEqual(d['total'], 2)
        self.assertEqual({c['type']: c['count'] for c in d['categories']},
                         {'ipv4-addr': 1, 'domain-name': 1})
        self.assertEqual([s['source'] for s in d['by_source']], ['otx'])
        self.assertEqual(d['sources'], 1)
        self.assertEqual(d['served'], 2)
        self.assertEqual(len(d['confidence']['histogram']), 5)
        self.assertIn('otx', [l['label'] for l in d['top_labels']])
        self.assertEqual(d['labels_total'], 3)   # otx, pulseA, otx

        self.assertEqual(
            self.client.get('/ui/stats?scope=manual', headers=self.auth).get_json()['total'], 1)
        all_stats = self.client.get('/ui/stats?scope=all', headers=self.auth).get_json()
        self.assertEqual(all_stats['total'], 3)
        self.assertEqual(all_stats['sources'], 2)

    def test_bad_scope_and_auth(self):
        self.assertEqual(
            self.client.get('/ui/stats?scope=nope', headers=self.auth).status_code, 400)
        self.assertEqual(self.client.get('/ui/stats').status_code, 401)

    def test_withheld_reason_reported(self):
        ingest_objects([
            {'id': 'ipv4-addr--10-0-0-5', 'type': 'ipv4-addr',
             'object': {'ipv4-addr': {'value': '10.0.0.5'}},
             'labels': ['otx'], 'confidence': 70},
        ], mode='merge', source='otx')
        d = self.client.get('/ui/stats?scope=community', headers=self.auth).get_json()
        self.assertEqual(d['withheld'], 1)
        self.assertEqual(d['served'], 0)
        self.assertEqual(d['withheld_reasons'][0]['reason'], 'private-ip')


class TestPullLimit(unittest.TestCase):
    """POST /community/pull/<name> — one-shot object limiter."""

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        app.config['TAXII_AUTH'] = {'username': 'u', 'password': 'p'}
        # Community pull is a feed-wide write: needs the admin principal.
        app.config['TAXII_ADMIN_AUTH'] = {'username': 'u', 'password': 'p'}
        self.auth = {'X-Taxii-Username': 'u', 'X-Taxii-Password': 'p'}

    def test_fetch_objects_honors_limit(self):
        import server
        p = server.TaxiiPuller({'name': 'c1', 'base_url': 'https://h/', 'collection': 'c'})
        p._get_taxii = lambda path, since=None: {
            'type': 'bundle',
            'objects': [{'type': 'ipv4-addr', 'value': f'1.1.1.{i}'} for i in range(10)]}
        self.assertEqual(len(p._fetch_objects(None, limit=3)), 3)   # capped by limit
        self.assertEqual(len(p._fetch_objects(None)), 10)          # default cap (5000)

    def test_endpoint_validates_limit(self):
        for bad in ('0', '-1', 'abc'):
            r = self.client.post('/community/pull/otx?limit=' + bad,
                                 headers=self.auth)
            self.assertEqual(r.status_code, 400, bad)              # no network hit

    def test_endpoint_passes_limit_to_puller(self):
        import server
        seen = {}

        class FakePuller:
            name = 'fakepuller'

            def pull_now(self, limit=None):
                seen['limit'] = limit
                return {'name': self.name, 'added': 0, 'error': None}

        server.taxii_pullers.append(FakePuller())
        try:
            r = self.client.post('/community/pull/fakepuller',
                                 json={'limit': 7}, headers=self.auth)
            self.assertEqual(r.status_code, 200)
            self.assertEqual(seen['limit'], 7)
            self.assertEqual(r.get_json()['limit'], 7)
        finally:
            server.taxii_pullers.pop()


class TestCollectionsRbac(unittest.TestCase):
    """Multi-collection support: per-collection credentials, isolated
    Get Objects, collection-scoped data endpoints, and the legacy
    single-collection default.
    """

    _GLOBALS = ('TAXII_COLLECTIONS', 'TAXII_COLLECTION_MAP',
                'TAXII_COLLECTION_ID', 'TAXII_COLLECTION_TITLE')

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        self._saved = {k: getattr(_server, k) for k in self._GLOBALS}
        cc = _server.CollectionConfig
        self.coll_a = cc(id='alpha', title='Alpha Feed',
                         description='first',
                         username='a_user', password='a_pass')
        self.coll_b = cc(id='beta', title='Beta Feed',
                         description='second',
                         username='b_user', password='b_pass')
        _server.TAXII_COLLECTIONS = [self.coll_a, self.coll_b]
        _server.TAXII_COLLECTION_MAP = {c.id: c for c in _server.TAXII_COLLECTIONS}
        _server.TAXII_COLLECTION_ID = 'alpha'
        _server.TAXII_COLLECTION_TITLE = 'Alpha Feed'
        # NO global credential override: only per-collection creds exist.
        app.config['TAXII_AUTH'] = {'username': '', 'password': ''}
        # Writes are the dashboard/admin's job; these creds stand in for it.
        app.config['TAXII_ADMIN_AUTH'] = {'username': 'admin', 'password': 'admin'}
        self.admin_auth = {'X-Taxii-Username': 'admin', 'X-Taxii-Password': 'admin'}
        self.auth_a = {'X-Taxii-Username': 'a_user', 'X-Taxii-Password': 'a_pass'}
        self.auth_b = {'X-Taxii-Username': 'b_user', 'X-Taxii-Password': 'b_pass'}

    def tearDown(self):
        for k, v in self._saved.items():
            setattr(_server, k, v)
        app.config.pop('TAXII_ADMIN_AUTH', None)

    def _ingest(self, objs, collection, auth=None):
        # Ingest targets a collection but is still a write: the read-only
        # collection credentials cannot do it, the admin principal can.
        return self.client.post(
            '/feed/ingest',
            json={'stix_objects': objs, 'collection': collection},
            headers=self.admin_auth)

    def test_collection_credentials_are_read_only(self):
        """A consumer credential can read, never write."""
        r = self._ingest([_manual_obj('1.1.1.1')], 'alpha')
        self.assertEqual(r.status_code, 200)
        # ... but the SAME principal cannot ingest its own collection ...
        r = self.client.post(
            '/feed/ingest',
            json={'stix_objects': [_manual_obj('9.9.9.9')], 'collection': 'alpha'},
            headers=self.auth_a)
        self.assertEqual(r.status_code, 401)
        # ... nor purge it, revoke, or trigger a community pull.
        r = self.client.delete('/feed/purge?collection=alpha', headers=self.auth_a)
        self.assertEqual(r.status_code, 401)
        r = self.client.post('/objects/ipv4-addr--x/revoke', json={},
                             headers=self.auth_a)
        self.assertEqual(r.status_code, 401)
        r = self.client.post('/community/pull/otx', headers=self.auth_a)
        self.assertEqual(r.status_code, 401)
        # The read it IS allowed to do still works.
        r = self.client.get('/taxii2/collections/alpha/objects/',
                            headers=self.auth_a)
        self.assertEqual(r.status_code, 200)

    def test_each_principal_sees_only_its_collection(self):
        self._ingest([_manual_obj('1.1.1.1')], 'alpha', self.auth_a)
        self._ingest([_manual_obj('2.2.2.2')], 'beta', self.auth_b)
        # alpha creds (collection A) list only A
        r = self.client.get('/taxii2/collections/', headers=self.auth_a)
        ids = [c['id'] for c in r.get_json()['collections']]
        self.assertEqual(ids, ['alpha'])
        # beta creds list only B
        r = self.client.get('/taxii2/collections/', headers=self.auth_b)
        ids = [c['id'] for c in r.get_json()['collections']]
        self.assertEqual(ids, ['beta'])

    def test_get_objects_isolated_per_collection(self):
        self._ingest([_manual_obj('1.1.1.1')], 'alpha', self.auth_a)
        self._ingest([_manual_obj('2.2.2.2')], 'beta', self.auth_b)
        # A's creds cannot read B
        r = self.client.get('/taxii2/collections/beta/objects/',
                            headers=self.auth_a)
        self.assertEqual(r.status_code, 401)
        # A's creds read only A's object
        r = self.client.get('/taxii2/collections/alpha/objects/',
                            headers=self.auth_a)
        objs = r.get_json()['content']['content']['objects']
        self.assertEqual(len(objs), 1)
        self.assertEqual(served_value(objs[0]), '1.1.1.1')

    def test_unknown_collection_404_and_bad_creds_401(self):
        self._ingest([_manual_obj('1.1.1.1')], 'alpha', self.auth_a)
        r = self.client.get('/taxii2/collections/nope/objects/',
                            headers=self.auth_a)
        self.assertEqual(r.status_code, 404)
        r = self.client.get('/taxii2/collections/nope/',
                            headers=self.auth_a)
        self.assertEqual(r.status_code, 404)
        # wrong creds entirely
        bad = {'X-Taxii-Username': 'x', 'X-Taxii-Password': 'y'}
        r = self.client.get('/taxii2/collections/alpha/objects/', headers=bad)
        self.assertEqual(r.status_code, 401)

    def test_data_endpoints_collection_scoping(self):
        self._ingest([_manual_obj('1.1.1.1')], 'alpha', self.auth_a)
        self._ingest([_manual_obj('2.2.2.2')], 'beta', self.auth_b)
        # /objects defaults to the primary collection (alpha)
        r = self.client.get('/objects', headers=self.auth_a)
        vals = {o['value'] for o in r.get_json()['objects']}
        self.assertEqual(vals, {'1.1.1.1'})
        # ... unless ?collection= says otherwise
        r = self.client.get('/objects?collection=beta', headers=self.auth_b)
        vals = {o['value'] for o in r.get_json()['objects']}
        self.assertEqual(vals, {'2.2.2.2'})
        # unknown collection -> 400
        r = self.client.get('/objects?collection=nope', headers=self.auth_a)
        self.assertEqual(r.status_code, 400)
        # ingest with unknown collection -> 400
        r = self._ingest([_manual_obj('3.3.3.3')], 'nope', self.auth_a)
        self.assertEqual(r.status_code, 400)

    def test_replace_ingest_is_per_collection(self):
        # a replace-save of beta must not touch alpha's manual rows
        self._ingest([_manual_obj('1.1.1.1')], 'alpha', self.auth_a)
        self._ingest([_manual_obj('2.2.2.2')], 'beta', self.auth_b)
        r = self.client.post(
            '/feed/ingest',
            json={'stix_objects': [_manual_obj('9.9.9.9')],
                  'collection': 'beta', 'mode': 'replace'},
            headers=self.admin_auth)
        self.assertEqual(r.status_code, 200)
        r = self.client.get('/objects', headers=self.auth_a)
        self.assertEqual({o['value'] for o in r.get_json()['objects']},
                         {'1.1.1.1'})
        r = self.client.get('/objects?collection=beta', headers=self.auth_b)
        self.assertEqual({o['value'] for o in r.get_json()['objects']},
                         {'9.9.9.9'})

    def test_purge_scoped_to_collection(self):
        self._ingest([_manual_obj('1.1.1.1')], 'alpha', self.auth_a)
        self._ingest([_manual_obj('2.2.2.2')], 'beta', self.auth_b)
        r = self.client.delete('/feed/purge?collection=beta',
                               headers=self.admin_auth)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['purged'], 1)
        r = self.client.get('/objects', headers=self.auth_a)
        self.assertEqual({o['value'] for o in r.get_json()['objects']},
                         {'1.1.1.1'})

    def test_ui_collections_endpoint(self):
        r = self.client.get('/ui/collections', headers=self.auth_a)
        self.assertEqual(r.status_code, 200)
        ids = [c['id'] for c in r.get_json()['collections']]
        self.assertEqual(ids, ['alpha', 'beta'])
        self.assertEqual(r.get_json()['primary'], 'alpha')

    def test_single_collection_compat_uses_primary(self):
        # With the registry pointing at one collection, the API behaves like
        # the legacy single-collection server.
        _server.TAXII_COLLECTIONS = [self.coll_a]
        _server.TAXII_COLLECTION_MAP = {'alpha': self.coll_a}
        self._ingest([_manual_obj('1.1.1.1')], 'alpha', self.auth_a)
        r = self.client.get('/taxii2/collections/', headers=self.auth_a)
        self.assertEqual(r.get_json()['meta']['count'], 1)
        r = self.client.get('/objects', headers=self.auth_a)
        self.assertEqual(r.get_json()['count'], 1)


class _CollectionRegistryTestCase(unittest.TestCase):
    """Shared plumbing for tests that swap the collection registry.

    Restores the module-level registry (and the two write-target constants)
    on teardown, and keeps the per-IP auth-failure throttle from leaking
    between tests — these tests deliberately provoke 401s.
    """

    _GLOBALS = ('TAXII_COLLECTIONS', 'TAXII_COLLECTION_MAP',
                'TAXII_COLLECTION_ID', 'TAXII_COLLECTION_TITLE',
                'DEFAULT_MANUAL_COLLECTION_ID',
                'DEFAULT_COMMUNITY_COLLECTION_ID')

    def _save_registry(self):
        self._saved = {k: getattr(_server, k) for k in self._GLOBALS}
        _server._auth_failures.clear()
        _server._api_key_hits.clear()

    def _restore_registry(self):
        for k, v in self._saved.items():
            setattr(_server, k, v)
        app.config.pop('TAXII_ADMIN_AUTH', None)
        _server._auth_failures.clear()
        _server._api_key_hits.clear()

    def _set_registry(self, collections):
        """Install ``collections`` as the registry, aggregate first."""
        _server.TAXII_COLLECTIONS = list(collections)
        _server.TAXII_COLLECTION_MAP = {c.id: c for c in collections}
        _server.TAXII_COLLECTION_ID = collections[0].id
        _server.TAXII_COLLECTION_TITLE = collections[0].title


class TestAggregateCollections(_CollectionRegistryTestCase):
    """Segregated collections: gt_intel (manual), comm_intel (community) and
    threat_intel — the read-only aggregate that serves their union.
    """

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        self._save_registry()
        cc = _server.CollectionConfig
        self.gt = cc(id='gt_intel', title='Ground truth', description='manual',
                     username='gt_user', password='gt_pass')
        self.comm = cc(id='comm_intel', title='Community', description='pulled',
                       username='comm_user', password='comm_pass')
        self.other = cc(id='sandbox', title='Sandbox', description='not a member',
                        username='s_user', password='s_pass')
        self.union = cc(id='threat_intel', title='Union', description='union',
                        username='union_user', password='union_pass',
                        members=['gt_intel', 'comm_intel'])
        self._set_registry([self.union, self.gt, self.comm, self.other])
        _server.DEFAULT_MANUAL_COLLECTION_ID = 'gt_intel'
        _server.DEFAULT_COMMUNITY_COLLECTION_ID = 'comm_intel'
        # NO global credential override: only per-collection creds exist.
        app.config['TAXII_AUTH'] = {'username': '', 'password': ''}
        app.config['TAXII_ADMIN_AUTH'] = {'username': 'admin', 'password': 'admin'}
        self.admin_auth = {'X-Taxii-Username': 'admin', 'X-Taxii-Password': 'admin'}
        self.union_auth = {'X-Taxii-Username': 'union_user',
                           'X-Taxii-Password': 'union_pass'}
        self.gt_auth = {'X-Taxii-Username': 'gt_user', 'X-Taxii-Password': 'gt_pass'}

    def tearDown(self):
        self._restore_registry()

    def _ingest(self, objs, collection=None, headers=None):
        body = {'stix_objects': objs, 'mode': 'merge'}
        if collection:
            body['collection'] = collection
        return self.client.post('/feed/ingest', json=body,
                                headers=headers or self.admin_auth)

    def _community(self, objs, collection='comm_intel'):
        """Write as a community puller would (source != 'manual')."""
        return _server.ingest_objects(objs, mode='merge', source='otx',
                                      collection=collection)

    def _placements(self):
        session = _server.create_session()
        try:
            return {r.stix_id: r.collection_id
                    for r in session.query(_server.STIXObject).all()}
        finally:
            session.close()

    def _taxii_values(self, collection, headers):
        r = self.client.get(f'/taxii2/collections/{collection}/objects/',
                            headers=headers)
        self.assertEqual(r.status_code, 200)
        return {served_value(o)
                for o in r.get_json()['content']['content']['objects']}

    # --- read path ------------------------------------------------------

    def test_aggregate_serves_the_union_of_its_members(self):
        self._ingest([_manual_obj('1.1.1.1')])                 # -> gt_intel
        self._community([_otx_obj('8.8.8.8')])                 # -> comm_intel
        self.assertEqual(self._placements(), {
            'ipv4-addr--1-1-1-1': 'gt_intel',
            'ipv4-addr--8-8-8-8': 'comm_intel',
        })
        self.assertEqual(self._taxii_values('threat_intel', self.union_auth),
                         {'1.1.1.1', '8.8.8.8'})
        # A member is still its own collection, not the union.
        self.assertEqual(self._taxii_values('gt_intel', self.gt_auth),
                         {'1.1.1.1'})

    def test_members_are_isolated_from_each_other(self):
        self._ingest([_manual_obj('1.1.1.1')])
        self._community([_otx_obj('8.8.8.8')])
        # The aggregate credential reads the union ...
        self.assertEqual(self._taxii_values('threat_intel', self.union_auth),
                         {'1.1.1.1', '8.8.8.8'})
        # ... but it grants no direct access to a member ...
        r = self.client.get('/taxii2/collections/gt_intel/objects/',
                            headers=self.union_auth)
        self.assertEqual(r.status_code, 401)
        # ... and a member credential cannot read a sibling, nor the union.
        r = self.client.get('/taxii2/collections/comm_intel/objects/',
                            headers=self.gt_auth)
        self.assertEqual(r.status_code, 401)
        r = self.client.get('/taxii2/collections/threat_intel/objects/',
                            headers=self.gt_auth)
        self.assertEqual(r.status_code, 401)

    def test_collection_listing_respects_the_scope(self):
        r = self.client.get('/taxii2/collections/', headers=self.union_auth)
        self.assertEqual([c['id'] for c in r.get_json()['collections']],
                         ['threat_intel'])
        r = self.client.get('/taxii2/collections/', headers=self.gt_auth)
        self.assertEqual([c['id'] for c in r.get_json()['collections']],
                         ['gt_intel'])

    def test_collection_resource_advertises_members(self):
        r = self.client.get('/taxii2/collections/threat_intel/',
                            headers=self.union_auth)
        self.assertEqual(r.status_code, 200)
        body = r.get_json()
        self.assertEqual(body['meta']['members'], ['gt_intel', 'comm_intel'])
        self.assertFalse(body['can_write'])
        # A leaf advertises no members (it is nobody's union).
        r = self.client.get('/taxii2/collections/gt_intel/', headers=self.gt_auth)
        self.assertNotIn('members', r.get_json()['meta'])

    def test_dashboard_defaults_and_column_reads_cover_the_union(self):
        self._ingest([_manual_obj('1.1.1.1')])
        self._community([_otx_obj('8.8.8.8')])
        # /objects and /ui/stats default to the primary (the union) ...
        r = self.client.get('/objects', headers=self.admin_auth)
        self.assertEqual({o['value'] for o in r.get_json()['objects']},
                         {'1.1.1.1', '8.8.8.8'})
        r = self.client.get('/ui/stats?scope=all', headers=self.admin_auth)
        self.assertEqual(r.get_json()['total'], 2)
        # ... and ?collection=<member> narrows it back down.
        r = self.client.get('/objects?collection=comm_intel', headers=self.admin_auth)
        self.assertEqual({o['value'] for o in r.get_json()['objects']}, {'8.8.8.8'})
        r = self.client.get('/feed', headers=self.admin_auth)
        self.assertEqual(r.status_code, 200)
        self.assertIn(b'1.1.1.1', r.data)
        self.assertIn(b'8.8.8.8', r.data)

    # --- write path -----------------------------------------------------

    def test_ingest_defaults_to_the_manual_collection(self):
        r = self._ingest([_manual_obj('1.1.1.1')])
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['collection'], 'gt_intel')
        self.assertEqual(self._placements(), {'ipv4-addr--1-1-1-1': 'gt_intel'})

    def test_write_to_the_aggregate_is_rejected(self):
        r = self._ingest([_manual_obj('1.1.1.1')], collection='threat_intel')
        self.assertEqual(r.status_code, 400)
        self.assertIn('aggregate', r.get_json()['error'])
        self.assertEqual(self._placements(), {})

    def test_csv_import_to_the_aggregate_is_rejected_and_defaults_to_manual(self):
        payload = 'ip\n1.1.1.1\n'
        r = self.client.post(
            '/feed/import-csv',
            data={'file': (io.BytesIO(payload.encode()), 'gt.csv'),
                  'collection': 'threat_intel'},
            headers=self.admin_auth)
        self.assertEqual(r.status_code, 400)
        r = self.client.post(
            '/feed/import-csv',
            data={'file': (io.BytesIO(payload.encode()), 'gt.csv')},
            headers=self.admin_auth)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self._placements(), {'ipv4-addr--1-1-1-1': 'gt_intel'})

    def test_replace_ingest_into_a_member_leaves_its_sibling_alone(self):
        self._ingest([_manual_obj('1.1.1.1')], collection='gt_intel')
        self._community([_otx_obj('8.8.8.8')])
        r = self.client.post(
            '/feed/ingest',
            json={'stix_objects': [_manual_obj('9.9.9.9')],
                  'collection': 'gt_intel', 'mode': 'replace'},
            headers=self.admin_auth)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self._placements(), {
            'ipv4-addr--9-9-9-9': 'gt_intel',
            'ipv4-addr--8-8-8-8': 'comm_intel',
        })

    def test_purging_the_aggregate_clears_members_only(self):
        self._ingest([_manual_obj('1.1.1.1')])
        self._community([_otx_obj('8.8.8.8')])
        self._community([_otx_obj('7.7.7.7')], collection='sandbox')
        r = self.client.delete('/feed/purge?collection=threat_intel',
                               headers=self.admin_auth)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['purged'], 2)
        # The non-member collection survives (unlike a collection-less purge).
        self.assertEqual(self._placements(), {'ipv4-addr--7-7-7-7': 'sandbox'})

    def test_purge_without_a_collection_still_wipes_everything(self):
        self._ingest([_manual_obj('1.1.1.1')])
        self._community([_otx_obj('7.7.7.7')], collection='sandbox')
        r = self.client.delete('/feed/purge', headers=self.admin_auth)
        self.assertEqual(r.get_json()['purged'], 2)
        self.assertEqual(self._placements(), {})

    # --- community ingest routing ---------------------------------------

    def test_pullers_target_the_configured_ingest_collection(self):
        p = _server.OtxPoller({'base_url': 'https://x/', 'enabled': False})
        self.assertEqual(p.ingest_collection, 'comm_intel')
        t = _server.TaxiiPuller({'name': 't', 'base_url': 'https://x/',
                                 'collection': 'remote-id'})
        # `collection` still means the UPSTREAM collection ...
        self.assertEqual(t.collection, 'remote-id')
        # ... while the local target defaults to the community leaf.
        self.assertEqual(t.ingest_collection, 'comm_intel')

    def test_ingest_collection_refuses_the_aggregate(self):
        # Pointing a poller at the union must not file community data there.
        t = _server.TaxiiPuller({'name': 't', 'base_url': 'https://x/',
                                 'collection': 'remote-id',
                                 'ingest_collection': 'threat_intel'})
        self.assertEqual(t.ingest_collection, 'comm_intel')
        # ... nor at an unknown collection.
        t = _server.TaxiiPuller({'name': 't', 'base_url': 'https://x/',
                                 'collection': 'remote-id',
                                 'ingest_collection': 'nope'})
        self.assertEqual(t.ingest_collection, 'comm_intel')

    # --- migration ------------------------------------------------------

    def _legacy_row(self, stix_id, source):
        session = _server.create_session()
        try:
            session.add(_server.STIXObject(
                id=stix_id, stix_id=stix_id, object_type='ipv4-addr',
                collection_id=_server.DEFAULT_COLLECTION_ID, source=source))
            session.commit()
        finally:
            session.close()

    def test_migration_rehomes_legacy_rows_and_is_idempotent(self):
        self._legacy_row('ipv4-addr--9-9-9-9', 'manual')
        self._legacy_row('ipv4-addr--8-8-8-8', 'otx')
        self._legacy_row('ipv4-addr--6-6-6-6', None)
        self.assertEqual(_server._rehome_aggregate_rows(), 3)
        self.assertEqual(self._placements(), {
            'ipv4-addr--9-9-9-9': 'gt_intel',    # manual
            'ipv4-addr--8-8-8-8': 'comm_intel',  # pulled
            'ipv4-addr--6-6-6-6': 'gt_intel',    # unknown provenance
        })
        # Idempotent: nothing carries the legacy id any more.
        self.assertEqual(_server._rehome_aggregate_rows(), 0)

    def test_migration_is_a_noop_when_the_feature_is_off(self):
        # With the legacy id still a collection of its own there is nothing to
        # re-home: its rows are exactly where they belong.
        legacy = _server.CollectionConfig(
            id=_server.DEFAULT_COLLECTION_ID, title='Legacy', description='',
            username='', password='')
        self._set_registry([legacy])
        self._legacy_row('ipv4-addr--9-9-9-9', 'manual')
        self.assertEqual(_server._rehome_aggregate_rows(), 0)
        self.assertEqual(self._placements(),
                         {'ipv4-addr--9-9-9-9': _server.DEFAULT_COLLECTION_ID})

    # --- configuration --------------------------------------------------

    def test_members_are_normalised(self):
        cfg = _server.CollectionConfig(
            id='a', title='', description='', username='', password='',
            members=['b', '', None, 'b', 'a', 'c'])
        # blanks dropped, duplicates collapsed, self-reference removed
        self.assertEqual(cfg.members, ['b', 'c'])
        self.assertTrue(cfg.is_aggregate)
        self.assertFalse(self.gt.is_aggregate)

    def test_unknown_and_nested_members_are_dropped(self):
        cc = _server.CollectionConfig
        inner = cc(id='inner', title='', description='', username='',
                   password='', members=['leaf'])
        outer = cc(id='outer', title='', description='', username='',
                   password='', members=['inner', 'leaf', 'ghost'])
        leaf = cc(id='leaf', title='', description='', username='', password='')
        self._set_registry([outer, inner, leaf])
        _server._validate_collection_members()
        self.assertEqual(outer.members, ['leaf'])   # nested + unknown dropped
        self.assertEqual(inner.members, ['leaf'])

    def test_aggregate_with_no_valid_member_is_a_startup_error(self):
        cc = _server.CollectionConfig
        broken = cc(id='broken', title='', description='', username='',
                    password='', members=['ghost'])
        self._set_registry([broken])
        with self.assertRaises(ValueError):
            _server._validate_collection_members()

    def test_config_parses_members_from_the_list_form(self):
        cfgs = _server._build_collection_configs({'collections': [
            {'id': 'union', 'members': ['a', 'b']},
            {'id': 'a'}, {'id': 'b'},
        ]})
        self.assertEqual([c.id for c in cfgs], ['union', 'a', 'b'])
        self.assertEqual(cfgs[0].members, ['a', 'b'])
        self.assertEqual(cfgs[1].members, [])
        # The legacy single-collection form yields no members.
        legacy = _server._build_collection_configs({'collection_id': 'solo'})
        self.assertEqual(legacy[0].members, [])
        self.assertEqual(legacy[0].id, 'solo')

    def test_ui_collections_reports_members_and_write_targets(self):
        r = self.client.get('/ui/collections', headers=self.admin_auth)
        self.assertEqual(r.status_code, 200)
        body = r.get_json()
        by_id = {c['id']: c for c in body['collections']}
        self.assertTrue(by_id['threat_intel']['aggregate'])
        self.assertEqual(by_id['threat_intel']['members'],
                         ['gt_intel', 'comm_intel'])
        self.assertFalse(by_id['gt_intel']['aggregate'])
        # Writes are routed without the dashboard hardcoding ids.
        self.assertEqual(body['primary'], 'threat_intel')
        self.assertEqual(body['manual'], 'gt_intel')
        self.assertEqual(body['community'], 'comm_intel')


class TestApiKeys(_CollectionRegistryTestCase):
    """Dashboard-issued, collection-scoped API keys."""

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        self._save_registry()
        cc = _server.CollectionConfig
        self.alpha = cc(id='alpha', title='Alpha', description='',
                        username='a_user', password='a_pass')
        self.beta = cc(id='beta', title='Beta', description='',
                       username='b_user', password='b_pass')
        self._set_registry([self.alpha, self.beta])
        app.config['TAXII_AUTH'] = {'username': '', 'password': ''}
        app.config['TAXII_ADMIN_AUTH'] = {'username': 'admin', 'password': 'admin'}
        self.admin_auth = {'X-Taxii-Username': 'admin', 'X-Taxii-Password': 'admin'}

    def tearDown(self):
        self._restore_registry()

    def _create(self, label='key', collections=('alpha',), types=None, **extra):
        # `types` is a required scope now; default it to every stored type so
        # tests that don't care about type filtering behave as before.
        body = {'label': label, 'collections': list(collections),
                'types': list(types) if types is not None
                         else sorted(_server._SUPPORTED_STIX_TYPES)}
        body.update(extra)
        return self.client.post('/ui/keys', json=body, headers=self.admin_auth)

    def _basic(self, credentials):
        """HTTP Basic headers for ``user:password`` — what Vision One speaks."""
        token = base64.b64encode(credentials.encode('utf-8')).decode('ascii')
        return {'Authorization': f'Basic {token}'}

    def _use(self, key, path='/taxii2/collections/', with_prefix=True):
        # A key rides in the password position; the username is its prefix
        # (or anything else — only the password is inspected).
        creds = f"{key.split('_')[1]}:{key}" if with_prefix else f"user:{key}"
        return self.client.get(path, headers=self._basic(creds))

    def test_key_is_returned_once_and_stored_hashed(self):
        r = self._create(label='vision one', collections=['alpha'])
        self.assertEqual(r.status_code, 201)
        body = r.get_json()
        key = body['key']
        self.assertTrue(key.startswith('taxii_'))
        self.assertEqual(body['collections'], ['alpha'])
        # Stored: only the prefix and a hash, never the secret.
        session = _server.create_session()
        try:
            row = session.query(_server.ApiKey).one()
        finally:
            session.close()
        self.assertNotIn(key, (row.secret_hash or ''))
        self.assertEqual(row.secret_hash,
                         _server._hash_api_key_secret(key))
        self.assertTrue(key.startswith('taxii_' + row.prefix + '_'))
        # The listing exposes the prefix, never the secret.
        r = self.client.get('/ui/keys', headers=self.admin_auth)
        listed = r.get_json()['keys']
        self.assertEqual(len(listed), 1)
        self.assertEqual(listed[0]['prefix'], row.prefix)
        self.assertNotIn('key', listed[0])
        self.assertNotIn('secret_hash', listed[0])

    def test_key_reads_only_its_own_scope(self):
        key = self._create(collections=['alpha']).get_json()['key']
        r = self._use(key)
        self.assertEqual([c['id'] for c in r.get_json()['collections']], ['alpha'])
        self.assertEqual(
            self._use(key, '/taxii2/collections/alpha/objects/').status_code, 200)
        # The sibling collection is not readable with a key scoped elsewhere.
        self.assertEqual(
            self._use(key, '/taxii2/collections/beta/objects/').status_code, 401)

    def test_key_never_writes(self):
        key = self._create(collections=['alpha']).get_json()['key']
        auth = self._basic(key)
        r = self.client.post('/feed/ingest',
                             json={'stix_objects': [_manual_obj('1.1.1.1')]},
                             headers=auth)
        self.assertEqual(r.status_code, 401)
        r = self.client.delete('/feed/purge?collection=alpha', headers=auth)
        self.assertEqual(r.status_code, 401)
        r = self.client.post('/community/pull/otx', headers=auth)
        self.assertEqual(r.status_code, 401)

    def test_revoked_key_is_rejected_and_others_survive(self):
        a = self._create(label='a', collections=['alpha']).get_json()
        b = self._create(label='b', collections=['beta']).get_json()
        self.assertEqual(self._use(a['key']).status_code, 200)
        r = self.client.post(f"/ui/keys/{a['id']}/revoke", headers=self.admin_auth)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self._use(a['key']).status_code, 401)
        # A different key is untouched.
        self.assertEqual(self._use(b['key']).status_code, 200)
        # Revoking twice reports no change rather than erroring.
        r = self.client.post(f"/ui/keys/{a['id']}/revoke", headers=self.admin_auth)
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.get_json()['changed'])

    def test_expired_key_is_rejected(self):
        created = self._create(expires_days=1).get_json()
        self.assertEqual(self._use(created['key']).status_code, 200)
        session = _server.create_session()
        try:
            row = session.query(_server.ApiKey).filter_by(id=created['id']).one()
            row.expires_at = datetime.utcnow() - timedelta(minutes=1)
            session.commit()
        finally:
            session.close()
        self.assertEqual(self._use(created['key']).status_code, 401)

    def test_rate_limit_is_per_key(self):
        limited = self._create(label='limited', collections=['alpha'],
                               rate_limit=2, rate_window=60).get_json()
        other = self._create(label='other', collections=['beta']).get_json()
        self.assertEqual(self._use(limited['key']).status_code, 200)
        self.assertEqual(self._use(limited['key']).status_code, 200)
        r = self._use(limited['key'])
        self.assertEqual(r.status_code, 429)
        self.assertIn('Retry-After', r.headers)
        self.assertGreaterEqual(int(r.headers['Retry-After']), 1)
        # A second key is unaffected by the first key's window.
        self.assertEqual(self._use(other['key']).status_code, 200)

    def test_key_creation_validates_its_input(self):
        self.assertEqual(self._create(label='').status_code, 400)
        self.assertEqual(self._create(collections=[]).status_code, 400)
        self.assertEqual(self._create(collections=['nope']).status_code, 400)
        self.assertEqual(self._create(expires_days=0).status_code, 400)
        self.assertEqual(self._create(rate_limit=-1).status_code, 400)
        # Object types are a required, strict allow-list.
        self.assertEqual(self._create(types=[]).status_code, 400)
        self.assertEqual(self._create(types=['nope']).status_code, 400)
        body = self.client.post('/ui/keys', json={
            'label': 'no-types', 'collections': ['alpha']},
            headers=self.admin_auth)
        self.assertEqual(body.status_code, 400)
        self.assertEqual(self._create(rate_window=0).status_code, 400)

    def test_key_management_requires_a_writer(self):
        # A read-only collection credential cannot mint or list keys.
        self.assertEqual(
            self.client.get('/ui/keys', headers={
                'X-Taxii-Username': 'a_user', 'X-Taxii-Password': 'a_pass',
            }).status_code, 401)
        self.assertEqual(self.client.get('/ui/keys').status_code, 401)
        self.assertEqual(self.client.post('/ui/keys', json={
            'label': 'x', 'collections': ['alpha']}).status_code, 401)

    def test_revoking_an_unknown_key_is_404(self):
        r = self.client.post('/ui/keys/nope/revoke', headers=self.admin_auth)
        self.assertEqual(r.status_code, 404)

    def test_aggregate_scoped_key_reads_the_union(self):
        # The Vision One case: one key that sees manual + community intel.
        cc = _server.CollectionConfig
        union = cc(id='threat_intel', title='Union', description='',
                   username='', password='', members=['alpha', 'beta'])
        self._set_registry([union, self.alpha, self.beta])
        _server.DEFAULT_MANUAL_COLLECTION_ID = 'alpha'
        _server.DEFAULT_COMMUNITY_COLLECTION_ID = 'beta'
        _server.ingest_objects([_manual_obj('1.1.1.1')], mode='merge',
                               source='manual', collection='alpha')
        _server.ingest_objects([_otx_obj('8.8.8.8')], mode='merge',
                               source='otx', collection='beta')
        key = self._create(collections=['threat_intel']).get_json()['key']
        r = self._use(key, '/taxii2/collections/threat_intel/objects/')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(
            {served_value(o) for o in r.get_json()['content']['content']['objects']},
            {'1.1.1.1', '8.8.8.8'})

    # --- per-response IOC ceiling (max_objects) ----------------------------

    def _seed(self, n, collection='alpha'):
        objs = [_manual_obj(f'10.0.0.{i}') for i in range(1, n + 1)]
        _server.ingest_objects(objs, mode='merge', source='manual',
                               collection=collection)
        return objs

    def _objects(self, key, path='/taxii2/collections/alpha/objects/'):
        r = self._use(key, path)
        self.assertEqual(r.status_code, 200)
        body = r.get_json()
        return body['content']['content']['objects'], body

    def test_key_max_objects_caps_each_response(self):
        """A key with a per-poll ceiling is never handed the whole feed."""
        self._seed(5)
        key = self._create(max_objects=2).get_json()['key']
        objs, body = self._objects(key)
        self.assertEqual(len(objs), 2)
        self.assertTrue(body['more'])
        self.assertIn('next', body)

    def test_key_max_objects_overrides_a_larger_client_limit(self):
        """A client asking for more than its key allows still gets the cap."""
        self._seed(5)
        key = self._create(max_objects=2).get_json()['key']
        objs, body = self._objects(
            key, '/taxii2/collections/alpha/objects/?limit=100')
        self.assertEqual(len(objs), 2)
        self.assertTrue(body['more'])

    def test_key_max_objects_pages_to_completion(self):
        """Following the cursor with a capped key loses nothing and repeats
        nothing — the property that makes the cap safe to rely on."""
        self._seed(5)
        key = self._create(max_objects=2).get_json()['key']
        seen, nxt, guard = [], None, 0
        while guard < 10:
            path = '/taxii2/collections/alpha/objects/'
            if nxt:
                path += '?next=' + urllib.parse.quote(nxt)
            objs, body = self._objects(key, path)
            seen.extend(served_value(o) for o in objs)
            if not body['more']:
                break
            nxt = body.get('next')
            guard += 1
        self.assertEqual(len(seen), 5)
        self.assertEqual(len(set(seen)), 5)

    def test_key_without_a_cap_is_not_limited(self):
        self._seed(5)
        key = self._create().get_json()['key']
        objs, body = self._objects(key)
        self.assertEqual(len(objs), 5)
        self.assertFalse(body['more'])

    def test_key_max_objects_is_validated_and_reported(self):
        self.assertEqual(self._create(label='neg', max_objects=-1).status_code, 400)
        self.assertEqual(
            self._create(label='junk', max_objects='lots').status_code, 400)
        self._create(label='capped', max_objects=7)
        keys = self.client.get('/ui/keys',
                               headers=self.admin_auth).get_json()['keys']
        capped = next(k for k in keys if k['label'] == 'capped')
        self.assertEqual(capped['max_objects'], 7)

    # --- per-key object-type allow-list ------------------------------------

    @staticmethod
    def _obj(stix_id, obj_type, props):
        return {'id': stix_id, 'type': obj_type, 'object': {obj_type: props},
                'labels': ['threat'], 'confidence': 90}

    def _seed_mixed(self, collection='alpha'):
        """An ipv4, a domain and a non-IOC (malware) object in the feed."""
        _server.ingest_objects([
            self._obj('ipv4-addr--10-0-0-1', 'ipv4-addr', {'value': '10.0.0.1'}),
            self._obj('domain-name--evil-example', 'domain-name',
                      {'value': 'evil.example'}),
            self._obj('malware--m1', 'malware', {'name': 'Zbot'}),
        ], mode='merge', source='manual', collection=collection)

    def _kinds(self, key, path='/taxii2/collections/alpha/objects/'):
        return {served_kind(o) for o in self._objects(key, path)[0]}

    def test_key_types_required_and_reported(self):
        # Required: absent/empty/unknown selections are all rejected.
        absent = self.client.post('/ui/keys', json={
            'label': 'k', 'collections': ['alpha']}, headers=self.admin_auth)
        self.assertEqual(absent.status_code, 400)
        self.assertEqual(self._create(types=[]).status_code, 400)
        self.assertEqual(self._create(types=['nope']).status_code, 400)
        # A valid selection is echoed back and listed on GET.
        body = self._create(label='typed',
                            types=['domain-name', 'ipv4-addr']).get_json()
        self.assertEqual(sorted(body['types']), ['domain-name', 'ipv4-addr'])
        listed = self.client.get('/ui/keys', headers=self.admin_auth).get_json()
        self.assertEqual(set(listed['types']), set(_server._SUPPORTED_STIX_TYPES))
        typed = next(k for k in listed['keys'] if k['label'] == 'typed')
        self.assertEqual(sorted(typed['types']), ['domain-name', 'ipv4-addr'])

    def test_key_serves_only_its_selected_types(self):
        self._seed_mixed()
        key = self._create(types=['ipv4-addr']).get_json()['key']
        self.assertEqual(self._kinds(key), {'ipv4-addr'})

    def test_key_types_are_a_strict_allow_list(self):
        """Non-IOC context objects are withheld unless the key lists them."""
        self._seed_mixed()
        ioc_only = self._create(
            label='ioc', types=['ipv4-addr', 'domain-name']).get_json()['key']
        self.assertEqual(self._kinds(ioc_only), {'ipv4-addr', 'domain-name'})
        with_malware = self._create(
            label='all',
            types=['ipv4-addr', 'domain-name', 'malware']).get_json()['key']
        self.assertIn('malware', self._kinds(with_malware))

    def test_indicator_selects_every_ioc_type(self):
        """'indicator' is an alias for all IOC values (they serve as indicators);
        it still does not unlock non-IOC objects."""
        self._seed_mixed()
        key = self._create(types=['indicator']).get_json()['key']
        self.assertEqual(self._kinds(key), {'ipv4-addr', 'domain-name'})

    def test_client_match_type_narrows_but_never_widens_the_key(self):
        self._seed_mixed()
        key = self._create(types=['domain-name']).get_json()['key']
        # The client asks for indicators; the key still only yields domains.
        self.assertEqual(
            self._kinds(key,
                        '/taxii2/collections/alpha/objects/?match[type]=indicator'),
            {'domain-name'})

    def test_legacy_row_without_types_serves_everything(self):
        """A pre-column (NULL types) key keeps serving all types."""
        self._seed_mixed()
        key = self._create(types=['ipv4-addr']).get_json()['key']
        session = _server.create_session()
        try:
            session.query(_server.ApiKey).delete()
            session.add(_server.ApiKey(
                id='legacy', prefix=key.split('_')[1],
                secret_hash=_server._hash_api_key_secret(key), enabled=True,
                collections='["alpha"]', types=None))
            session.commit()
        finally:
            session.close()
        self.assertEqual(self._kinds(key),
                         {'ipv4-addr', 'domain-name', 'malware'})


class TestSecurityHardening(unittest.TestCase):
    """Response headers, cookie flags, upload guards and auth throttling."""

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        app.config['TAXII_AUTH'] = {'username': 'u', 'password': 'p'}
        app.config['TAXII_ADMIN_AUTH'] = {'username': 'u', 'password': 'p'}
        self.auth = {'X-Taxii-Username': 'u', 'X-Taxii-Password': 'p'}
        self._ui_saved = (_server._ui_username, _server._ui_password)
        _server._auth_failures.clear()

    def tearDown(self):
        app.config.pop('TAXII_ADMIN_AUTH', None)
        _server._ui_username, _server._ui_password = self._ui_saved
        _server._auth_failures.clear()
        _server._AUTH_FAIL_LIMIT = 10

    def test_security_headers_on_every_response(self):
        for path in ('/', '/health'):
            r = self.client.get(path)
            self.assertEqual(r.status_code, 200)
            for header in ('Content-Security-Policy', 'X-Frame-Options',
                           'X-Content-Type-Options', 'Referrer-Policy',
                           'Permissions-Policy', 'Strict-Transport-Security'):
                self.assertIn(header, r.headers, f'{header} missing on {path}')
            self.assertEqual(r.headers['X-Frame-Options'], 'DENY')
            self.assertEqual(r.headers['X-Content-Type-Options'], 'nosniff')
            self.assertIn("frame-ancestors 'none'",
                          r.headers['Content-Security-Policy'])

    def test_login_cookie_is_hardened(self):
        _server._ui_username, _server._ui_password = 'ui', 'pw'
        r = self.client.post('/ui/login',
                             json={'username': 'ui', 'password': 'pw'})
        self.assertEqual(r.status_code, 200)
        cookie = r.headers.get('Set-Cookie', '')
        self.assertIn('HttpOnly', cookie)
        self.assertIn('Secure', cookie)
        self.assertIn('SameSite=Lax', cookie)

    def test_upload_over_cap_is_413(self):
        original = app.config['MAX_CONTENT_LENGTH']
        app.config['MAX_CONTENT_LENGTH'] = 1024
        try:
            big = b'ip\n' + b'1.2.3.4\n' * 500
            r = self.client.post(
                '/feed/import-csv',
                data={'file': (io.BytesIO(big), 'big.csv')},
                headers=self.auth)
            self.assertEqual(r.status_code, 413)
        finally:
            app.config['MAX_CONTENT_LENGTH'] = original

    def test_bad_workbook_returns_400_without_internals(self):
        r = self.client.post(
            '/feed/import-csv',
            data={'file': (io.BytesIO(b'not really a zip'), 'x.xlsx')},
            headers=self.auth)
        self.assertEqual(r.status_code, 400)
        body = r.get_data(as_text=True)
        for leak in ('Traceback', 'zipfile', 'openpyxl', '/app/', 'BadZipFile',
                     'not a zip', 'not really a zip'):
            self.assertNotIn(leak, body)

    def test_hostile_csv_values_are_stored_inert(self):
        """Formula / script / SQLi / traversal values are data, not code."""
        payload = (
            'ip,domain,indicator,label\n'
            '1.1.1.1,evil.example,=cmd|\' /C calc\'!A0,=HYPERLINK("http://x")\n'
            '2.2.2.2,evil2.example,<script>alert(1)</script>,x\n'
            '3.3.3.3,evil3.example,1;DROP TABLE stix_objects;--,y\n'
            '4.4.4.4,evil4.example,../../../../etc/passwd,z\n'
        )
        r = self.client.post(
            '/feed/import-csv',
            data={'file': (io.BytesIO(payload.encode()), 'hostile.csv')},
            headers=self.auth)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['imported'], 12)
        # Nothing executed: the table is intact and every value round-trips
        # verbatim as data...
        served = self.client.get(
            '/taxii2/collections/threat_intel/objects/',
            headers=self.auth).get_json()['content']['content']['objects']
        values = {served_value(o) for o in served}
        for literal in ("=cmd|' /C calc'!A0", '<script>alert(1)</script>',
                        '1;DROP TABLE stix_objects;--', '../../../../etc/passwd'):
            self.assertIn(literal, values)
        # ... and the dashboard JSON keeps the feed served, not dropped.
        self.assertEqual(self.client.get('/objects', headers=self.auth
                                         ).get_json()['count'], 12)

    def test_hostile_csv_headers_cannot_set_privileged_fields(self):
        """Columns named source/collection_id/stix_id/revoked are ignored."""
        payload = ('source,collection_id,stix_id,revoked,value\n'
                   'otx,premium,ipv4-addr--9-9-9-9,true,9.9.9.9\n')
        r = self.client.post(
            '/feed/import-csv',
            data={'file': (io.BytesIO(payload.encode()), 'h.csv')},
            headers=self.auth)
        self.assertEqual(r.status_code, 200)
        obj = self.client.get('/objects', headers=self.auth).get_json()['objects'][0]
        self.assertEqual(obj['value'], '9.9.9.9')
        self.assertEqual(obj['source'], 'manual')       # not 'otx'
        # a 'value:' column is a free-text indicator; the caller's stix_id
        # and collection_id columns carry no authority at all
        self.assertEqual(obj['id'], 'indicator--9-9-9-9')
        self.assertNotEqual(obj['id'], 'ipv4-addr--9-9-9-9')
        self.assertFalse(obj['revoked'])

    def test_xlsx_zip_bomb_is_bounded(self):
        """A highly compressible workbook stops at max_import_rows."""
        import zipfile
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
            z.writestr('xl/worksheets/sheet1.xml', '<!-- pad -->' * 200000)
        r = self.client.post(
            '/feed/import-csv',
            data={'file': (io.BytesIO(buf.getvalue()), 'bomb.xlsx')},
            headers=self.auth)
        # Either rejected outright (400) or truncated — never unbounded/OOM.
        self.assertIn(r.status_code, (400, 200))

    def test_objects_limit_and_truncation_flag(self):
        self.client.post('/feed/ingest', json={'stix_objects': [
            _manual_obj(f'203.0.113.{i}') for i in range(1, 6)]},
            headers=self.auth)
        d = self.client.get('/objects?limit=2', headers=self.auth).get_json()
        self.assertEqual(d['count'], 2)
        self.assertTrue(d['truncated'])
        d = self.client.get('/objects', headers=self.auth).get_json()
        self.assertEqual(d['count'], 5)
        self.assertFalse(d['truncated'])
        self.assertEqual(
            self.client.get('/objects?limit=0', headers=self.auth).status_code,
            400)

    def test_repeated_credentialed_failures_are_throttled(self):
        _server._AUTH_FAIL_LIMIT = 3
        bad = {'X-Taxii-Username': 'u', 'X-Taxii-Password': 'wrong'}
        codes = [self.client.get('/objects', headers=bad).status_code
                 for _ in range(4)]
        self.assertEqual(codes, [401, 401, 401, 429])
        # while throttled the credential is not even evaluated
        self.assertEqual(
            self.client.get('/objects', headers=self.auth).status_code, 429)
        _server._auth_failures.clear()
        self.assertEqual(
            self.client.get('/objects', headers=self.auth).status_code, 200)

    def test_anonymous_401s_are_not_counted(self):
        for _ in range(12):
            self.assertEqual(self.client.get('/objects').status_code, 401)
        self.assertEqual(_server._auth_failures, {})

    def test_sso_error_page_escapes_caller_input(self):
        with app.test_request_context('/ui/sso/callback'):
            resp = _server._ui_sso_fail(
                'Microsoft returned an error: <script>alert(1)</script>')
            body = resp.get_data(as_text=True)
        self.assertNotIn('<script>', body)
        self.assertIn('&lt;script&gt;', body)

    def test_unknown_collection_needs_credentials_to_be_distinguishable(self):
        # Unauthenticated: 401 whether or not the collection exists (no oracle).
        self.assertEqual(
            self.client.get('/taxii2/collections/nope/objects/').status_code,
            401)
        self.assertEqual(
            self.client.get('/taxii2/collections/threat_intel/objects/'
                            ).status_code, 401)
        # Authenticated: a real collection resolves, an unknown one 404s.
        self.assertEqual(
            self.client.get('/taxii2/collections/nope/objects/',
                            headers=self.auth).status_code, 404)
        self.assertEqual(
            self.client.get('/taxii2/collections/threat_intel/objects/',
                            headers=self.auth).status_code, 200)


class TestServedTimestamps(unittest.TestCase):
    """The feed is served with the STORED created/modified, not render-time now.

    Stamping ``now`` at render time made every object look newly modified on
    every poll: it broke ``added_after`` delta sync (a client's watermark no
    longer corresponded to storage) and forced consumers to re-process the
    whole feed each cycle.
    """

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        app.config['TAXII_AUTH'] = {'username': 'u', 'password': 'p'}
        app.config['TAXII_ADMIN_AUTH'] = {'username': 'u', 'password': 'p'}
        self.auth = {'X-Taxii-Username': 'u', 'X-Taxii-Password': 'p'}

    def tearDown(self):
        app.config.pop('TAXII_ADMIN_AUTH', None)

    def _served(self):
        r = self.client.get('/taxii2/collections/threat_intel/objects/',
                            headers=self.auth)
        self.assertEqual(r.status_code, 200)
        return r.get_json()['content']['content']['objects']

    def _ingest(self, value, mode=None):
        body = {
            'stix_objects': [{
                'id': f'ipv4-addr--{value.replace(".", "-")}',
                'type': 'ipv4-addr',
                'object': {'ipv4-addr': {'value': value}},
                'labels': ['t'], 'confidence': 50,
            }],
        }
        if mode:
            body['mode'] = mode
        r = self.client.post('/feed/ingest', json=body, headers=self.auth)
        self.assertEqual(r.status_code, 200)
        return body['stix_objects'][0]['id']

    def _backdate(self, sid, when):
        session = _server.create_session()
        try:
            row = session.query(_server.STIXObject).filter_by(stix_id=sid).one()
            row.created = when
            row.modified = when
            session.commit()
        finally:
            session.close()

    def test_served_timestamps_are_the_stored_ones_and_stay_stable(self):
        sid = self._ingest('9.9.9.9')
        old = datetime.utcnow() - timedelta(days=40)
        self._backdate(sid, old)
        expected = old.strftime('%Y-%m-%dT%H:%M:%SZ')

        first = self._served()[0]
        second = self._served()[0]
        # A render-time `now` would report today here, not 40 days ago.
        self.assertEqual(first['created'], expected)
        self.assertEqual(first['modified'], expected)
        # ...and would drift between polls.
        self.assertEqual(second['created'], first['created'])
        self.assertEqual(second['modified'], first['modified'])

    def test_merge_refresh_keeps_created_and_moves_modified(self):
        """created is immutable per object; modified advances on a change."""
        sid = self._ingest('5.5.5.5')
        before = self._served()[0]
        time.sleep(1.05)  # served timestamps are second-precision
        body = {
            'mode': 'merge',
            'stix_objects': [{
                'id': sid, 'type': 'ipv4-addr',
                'object': {'ipv4-addr': {'value': '5.5.5.5'}},
                'labels': ['t'], 'confidence': 90,
            }],
        }
        self.assertEqual(
            self.client.post('/feed/ingest', json=body,
                             headers=self.auth).status_code, 200)
        after = self._served()[0]
        self.assertEqual(after['created'], before['created'])
        self.assertGreater(after['modified'], before['modified'])

    def test_added_after_watermark_is_a_real_resume_point(self):
        """A watermark taken from a served `modified` filters on storage."""
        stale = self._ingest('1.1.1.1')
        self._backdate(stale, datetime.utcnow() - timedelta(days=40))
        self._ingest('2.2.2.2')
        watermark = self._served()[0]['modified']  # the backdated object
        r = self.client.get(
            '/taxii2/collections/threat_intel/objects/?added_after='
            + urllib.parse.quote(watermark),
            headers=self.auth)
        self.assertEqual(r.status_code, 200)
        values = {served_value(o)
                  for o in r.get_json()['content']['content']['objects']}
        self.assertIn('2.2.2.2', values)
        self.assertNotIn('1.1.1.1', values)  # older than the watermark


class TestFeedReadsFromTheDatabase(unittest.TestCase):
    """What is served comes from the DB — a stale cache cannot hide an object."""

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        app.config['TAXII_AUTH'] = {'username': 'u', 'password': 'p'}
        app.config['TAXII_ADMIN_AUTH'] = {'username': 'u', 'password': 'p'}
        self.auth = {'X-Taxii-Username': 'u', 'X-Taxii-Password': 'p'}

    def tearDown(self):
        app.config.pop('TAXII_ADMIN_AUTH', None)

    def test_poll_does_not_depend_on_the_memory_cache(self):
        r = self.client.post('/feed/ingest', json={'stix_objects': [{
            'id': 'ipv4-addr--7-7-7-7', 'type': 'ipv4-addr',
            'object': {'ipv4-addr': {'value': '7.7.7.7'}},
            'labels': ['t'], 'confidence': 50}]}, headers=self.auth)
        self.assertEqual(r.status_code, 200)
        # Simulate a cache that lost the entry — the old read path skipped
        # such objects silently, shortening pages and dropping data.
        with _server.memory_lock:
            _server.memory_store.clear()
        resp = self.client.get('/taxii2/collections/threat_intel/objects/',
                               headers=self.auth)
        self.assertEqual(resp.status_code, 200)
        values = {served_value(o)
                  for o in resp.get_json()['content']['content']['objects']}
        self.assertIn('7.7.7.7', values)


class TestPollPathIndexes(unittest.TestCase):
    """Get Objects filters/paginates on `modified`: the index must exist."""

    def test_modified_indexes_are_created(self):
        init_db()
        insp = sqlalchemy.inspect(_server.engine)
        names = {ix['name'] for ix in insp.get_indexes('stix_objects')}
        self.assertIn('ix_stix_objects_modified_id', names)
        self.assertIn('ix_stix_objects_coll_modified_id', names)


class TestThrottleTableBounded(unittest.TestCase):
    """The auth-failure table evicts the oldest keys; it never clears wholesale.

    A wholesale clear let a caller able to vary its throttle key wipe
    brute-force protection for every other client in a single request.
    """

    def setUp(self):
        self._saved_max = _server._AUTH_FAIL_MAX_KEYS
        _server._auth_failures.clear()

    def tearDown(self):
        _server._AUTH_FAIL_MAX_KEYS = self._saved_max
        _server._auth_failures.clear()

    def test_capacity_does_not_wipe_the_table(self):
        _server._AUTH_FAIL_MAX_KEYS = 10
        for i in range(10):
            _server._record_auth_failure(f'ip-{i}')
        self.assertEqual(len(_server._auth_failures), 10)
        _server._record_auth_failure('ip-10')  # tips over capacity
        self.assertGreater(len(_server._auth_failures), 1)  # not a clear()
        self.assertLessEqual(len(_server._auth_failures), 10)
        self.assertIn('ip-9', _server._auth_failures)

    def test_capacity_evicts_the_oldest_keys(self):
        _server._AUTH_FAIL_MAX_KEYS = 10
        for i in range(10):
            _server._record_auth_failure(f'ip-{i}')
        for i in range(10, 25):
            _server._record_auth_failure(f'ip-{i}')
        self.assertNotIn('ip-0', _server._auth_failures)
        self.assertIn('ip-24', _server._auth_failures)
        self.assertLessEqual(len(_server._auth_failures), 10)


class TestClientIpTrust(unittest.TestCase):
    """Proxy client-IP headers are honoured only from a trusted proxy."""

    def test_spoofed_header_from_an_untrusted_peer_is_ignored(self):
        with app.test_request_context(
                '/', environ_base={'REMOTE_ADDR': '203.0.113.9'},
                headers={'CF-Connecting-IP': '198.51.100.7'}):
            self.assertEqual(_server._client_ip(), '203.0.113.9')

    def test_header_from_a_trusted_proxy_is_honoured(self):
        with app.test_request_context(
                '/', environ_base={'REMOTE_ADDR': '127.0.0.1'},
                headers={'CF-Connecting-IP': '198.51.100.7'}):
            self.assertEqual(_server._client_ip(), '198.51.100.7')

    def test_first_forwarded_hop_is_used_from_a_trusted_proxy(self):
        with app.test_request_context(
                '/', environ_base={'REMOTE_ADDR': '192.168.1.5'},
                headers={'X-Forwarded-For': '198.51.100.7, 10.0.0.1'}):
            self.assertEqual(_server._client_ip(), '198.51.100.7')

    def test_no_trusted_proxies_configured_uses_the_peer(self):
        saved = _server._TRUSTED_PROXY_NETS
        _server._TRUSTED_PROXY_NETS = []
        try:
            with app.test_request_context(
                    '/', environ_base={'REMOTE_ADDR': '127.0.0.1'},
                    headers={'CF-Connecting-IP': '198.51.100.7'}):
                self.assertEqual(_server._client_ip(), '127.0.0.1')
        finally:
            _server._TRUSTED_PROXY_NETS = saved


class TestDashboardUsers(unittest.TestCase):
    """Dashboard multi-user logins with roles (viewer / editor / admin).

    Permissions are enforced server-side; the role is re-resolved from the
    users table on every request, so changes apply without a re-login.
    """

    ADMIN = ('root', 'root-password-123')
    PW = 'password-12345'

    def setUp(self):
        _reset_db()
        self._saved = (_server._ui_username, _server._ui_password)
        _server._ui_username, _server._ui_password = self.ADMIN
        self.admin = app.test_client()
        r = self.admin.post('/ui/login', json={
            'username': self.ADMIN[0], 'password': self.ADMIN[1]})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))

    def tearDown(self):
        _server._ui_username, _server._ui_password = self._saved

    def _mkuser(self, username, role):
        r = self.admin.post('/ui/users', json={
            'username': username, 'password': self.PW, 'role': role})
        self.assertEqual(r.status_code, 201, r.get_data(as_text=True))
        return r.get_json()

    def _login(self, username, password=None):
        client = app.test_client()
        r = client.post('/ui/login', json={
            'username': username, 'password': password or self.PW})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        return client

    def _ingest(self, client):
        return client.post('/feed/ingest',
                           json={'stix_objects': [_manual_obj('7.7.7.7')]})

    # --- provisioning -------------------------------------------------------

    def test_admin_creates_user_who_can_log_in(self):
        self._mkuser('alice', 'editor')
        session = self._login('alice').get('/ui/session').get_json()
        self.assertEqual(session['user'], 'alice')
        self.assertEqual(session['role'], 'editor')
        self.assertTrue(session['can_write'])
        self.assertFalse(session['can_admin'])

    def test_password_hash_is_never_returned(self):
        body = self._mkuser('bob', 'viewer')
        self.assertNotIn('password', body)
        self.assertNotIn('password_hash', body)

    def test_duplicate_username_rejected(self):
        self._mkuser('dup', 'viewer')
        self.assertEqual(self.admin.post('/ui/users', json={
            'username': 'dup', 'password': self.PW, 'role': 'viewer'}).status_code,
            409)

    def test_role_and_password_are_validated(self):
        self.assertEqual(self.admin.post('/ui/users', json={
            'username': 'x', 'password': self.PW, 'role': 'superuser'}
        ).status_code, 400)
        self.assertEqual(self.admin.post('/ui/users', json={
            'username': 'y', 'password': 'short', 'role': 'viewer'}
        ).status_code, 400)

    def test_static_credentials_remain_a_break_glass_admin(self):
        session = self.admin.get('/ui/session').get_json()
        self.assertEqual(session['role'], 'admin')
        self.assertTrue(session['can_admin'])

    # --- role enforcement ---------------------------------------------------

    def test_viewer_can_read_but_not_write(self):
        self._mkuser('vi', 'viewer')
        viewer = self._login('vi')
        self.assertEqual(viewer.get('/objects').status_code, 200)
        self.assertEqual(self._ingest(viewer).status_code, 401)
        self.assertEqual(viewer.post('/ui/keys', json={
            'label': 'x', 'collections': ['gt_intel']}).status_code, 401)

    def test_editor_can_write_but_not_manage_keys_or_users(self):
        self._mkuser('ed', 'editor')
        editor = self._login('ed')
        self.assertEqual(self._ingest(editor).status_code, 200)
        self.assertEqual(editor.post('/ui/keys', json={
            'label': 'x', 'collections': ['gt_intel']}).status_code, 401)
        self.assertEqual(editor.get('/ui/users').status_code, 401)

    def test_admin_can_manage_keys_and_users(self):
        self._mkuser('adm', 'admin')
        admin = self._login('adm')
        self.assertEqual(admin.get('/ui/users').status_code, 200)
        self.assertEqual(admin.post('/ui/keys', json={
            'label': 'k', 'collections': ['gt_intel'],
            'types': ['ipv4-addr']}).status_code, 201)

    # --- changes apply without a re-login -----------------------------------

    def test_disabling_a_user_ends_their_session_immediately(self):
        user = self._mkuser('temp', 'editor')
        temp = self._login('temp')
        self.assertEqual(temp.get('/objects').status_code, 200)
        self.assertEqual(self.admin.post(
            f"/ui/users/{user['id']}", json={'enabled': False}).status_code, 200)
        self.assertEqual(temp.get('/objects').status_code, 401)

    def test_deleting_a_user_ends_their_session(self):
        user = self._mkuser('gone', 'editor')
        gone = self._login('gone')
        self.assertEqual(self.admin.post(
            f"/ui/users/{user['id']}/delete").status_code, 200)
        self.assertEqual(gone.get('/objects').status_code, 401)

    def test_role_change_takes_effect_without_relogin(self):
        user = self._mkuser('promo', 'viewer')
        promo = self._login('promo')
        self.assertEqual(self._ingest(promo).status_code, 401)
        self.admin.post(f"/ui/users/{user['id']}", json={'role': 'editor'})
        self.assertEqual(self._ingest(promo).status_code, 200)

    # --- lockout guard ------------------------------------------------------

    def test_last_enabled_admin_cannot_be_removed(self):
        user = self._mkuser('only-admin', 'admin')
        self.assertEqual(self.admin.post(
            f"/ui/users/{user['id']}", json={'role': 'viewer'}).status_code, 400)
        self.assertEqual(self.admin.post(
            f"/ui/users/{user['id']}", json={'enabled': False}).status_code, 400)
        self.assertEqual(self.admin.post(
            f"/ui/users/{user['id']}/delete").status_code, 400)

    def test_a_second_admin_can_be_removed(self):
        first = self._mkuser('admin-a', 'admin')
        self._mkuser('admin-b', 'admin')
        self.assertEqual(self.admin.post(
            f"/ui/users/{first['id']}/delete").status_code, 200)

    # --- audit --------------------------------------------------------------

    def test_write_audit_line_names_the_actor(self):
        self._mkuser('audited', 'editor')
        editor = self._login('audited')
        with self.assertLogs('TAXII', level='INFO') as logs:
            self._ingest(editor)
        joined = '\n'.join(logs.output)
        self.assertIn('audited', joined)
        self.assertIn('editor', joined)


class TestAuditLog(unittest.TestCase):
    """The durable audit trail (GET /ui/audit) and its call sites.

    Every successful mutation and login appends a row; failed requests do not,
    and secrets never appear. The trail is admin-only.
    """

    ADMIN = ('root', 'root-password-123')
    PW = 'password-12345'

    def setUp(self):
        _reset_db()
        self._saved = (_server._ui_username, _server._ui_password)
        _server._ui_username, _server._ui_password = self.ADMIN
        self.admin = app.test_client()
        r = self.admin.post('/ui/login', json={
            'username': self.ADMIN[0], 'password': self.ADMIN[1]})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))

    def tearDown(self):
        _server._ui_username, _server._ui_password = self._saved

    # --- helpers -----------------------------------------------------------

    def _mkuser(self, username, role):
        r = self.admin.post('/ui/users', json={
            'username': username, 'password': self.PW, 'role': role})
        self.assertEqual(r.status_code, 201, r.get_data(as_text=True))
        return r.get_json()

    def _login(self, username, password=None):
        client = app.test_client()
        r = client.post('/ui/login', json={
            'username': username, 'password': password or self.PW})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        return client

    def _events(self, **filters):
        """Audit rows matching ``filters``, oldest first (direct DB read)."""
        session = _server.create_session()
        try:
            query = session.query(_server.AuditEvent)
            for key, value in filters.items():
                query = query.filter(getattr(_server.AuditEvent, key) == value)
            return query.order_by(_server.AuditEvent.created.asc(),
                                  _server.AuditEvent.id.asc()).all()
        finally:
            session.close()

    def _one(self, **filters):
        rows = self._events(**filters)
        self.assertEqual(len(rows), 1,
                         f'expected one {filters!r} row, got {len(rows)}')
        return rows[0]

    # --- access control ----------------------------------------------------

    def test_admin_can_read_the_log(self):
        r = self.admin.get('/ui/audit')
        self.assertEqual(r.status_code, 200)
        body = r.get_json()
        self.assertIn('events', body)
        self.assertEqual(body['limit'], 200)
        self.assertEqual(body['max_limit'], 1000)
        self.assertIn('feed.ingest', body['actions'])

    def test_non_admins_get_401(self):
        self._mkuser('ed', 'editor')
        self._mkuser('vi', 'viewer')
        self.assertEqual(self._login('ed').get('/ui/audit').status_code, 401)
        self.assertEqual(self._login('vi').get('/ui/audit').status_code, 401)
        self.assertEqual(app.test_client().get('/ui/audit').status_code, 401)

    # --- what gets recorded ------------------------------------------------

    def test_static_login_is_recorded(self):
        row = self._one(action='auth.login', actor='root')
        self.assertEqual(row.role, 'admin')
        self.assertEqual(row.via, 'ui_auth')     # static break-glass pair
        self.assertEqual(row.detail, 'local')

    def test_real_user_login_is_recorded(self):
        self._mkuser('alice', 'editor')
        self._login('alice')
        row = self._one(action='auth.login', actor='alice')
        self.assertEqual(row.role, 'editor')
        self.assertEqual(row.via, 'session')
        self.assertEqual(row.detail, 'local')

    def test_ingest_records_actor_role_via_target_and_count(self):
        self._mkuser('ed', 'editor')
        editor = self._login('ed')
        r = editor.post('/feed/ingest',
                        json={'stix_objects': [_manual_obj('7.7.7.7')]})
        self.assertEqual(r.status_code, 200)
        row = self._one(action='feed.ingest', actor='ed')
        self.assertEqual(row.role, 'editor')
        self.assertEqual(row.via, 'session')
        self.assertEqual(row.target, _server.DEFAULT_MANUAL_COLLECTION_ID)
        self.assertEqual(row.collection, _server.DEFAULT_MANUAL_COLLECTION_ID)
        self.assertEqual(row.count, 1)
        self.assertEqual(row.detail, 'mode=replace')
        self.assertTrue(row.ip)

    def test_revoke_records_the_object_and_source(self):
        self._mkuser('ed', 'editor')
        editor = self._login('ed')
        editor.post('/feed/ingest', json={'stix_objects': [_manual_obj('7.7.7.7')]})
        stix_id = _manual_obj('7.7.7.7')['id']
        r = editor.post(f'/objects/{stix_id}/revoke', json={'action': 'revoke'})
        self.assertEqual(r.status_code, 200)
        row = self._one(action='object.revoke', target=stix_id)
        self.assertEqual(row.actor, 'ed')
        self.assertEqual(row.collection, _server.DEFAULT_MANUAL_COLLECTION_ID)
        self.assertEqual(row.detail, 'manual')

    def test_delete_records_count_and_ids(self):
        self._mkuser('ed', 'editor')
        editor = self._login('ed')
        editor.post('/feed/ingest', json={'stix_objects': [_manual_obj('7.7.7.7')]})
        stix_id = _manual_obj('7.7.7.7')['id']
        r = editor.post('/feed/delete', json={'ids': [stix_id]})
        self.assertEqual(r.status_code, 200)
        row = self._one(action='feed.delete')
        self.assertEqual(row.count, 1)
        self.assertIn(stix_id, row.detail)

    def test_purge_records_the_collection(self):
        self._mkuser('ed', 'editor')
        editor = self._login('ed')
        editor.post('/feed/ingest', json={'stix_objects': [_manual_obj('7.7.7.7')]})
        r = editor.delete(
            f"/feed/purge?collection={_server.DEFAULT_MANUAL_COLLECTION_ID}")
        self.assertEqual(r.status_code, 200)
        row = self._one(action='feed.purge')
        self.assertEqual(row.target, _server.DEFAULT_MANUAL_COLLECTION_ID)
        self.assertEqual(row.count, 1)

    def test_key_create_records_prefix_never_the_secret(self):
        r = self.admin.post('/ui/keys', json={
            'label': 'ci-key',
            'collections': [_server.DEFAULT_MANUAL_COLLECTION_ID],
            'types': ['ipv4-addr']})
        self.assertEqual(r.status_code, 201)
        body = r.get_json()
        row = self._one(action='key.create')
        self.assertEqual(row.target, body['prefix'])
        self.assertNotIn(body['key'], row.detail or '')
        self.assertNotIn(body['key'], row.target or '')
        self.assertIn('ci-key', row.detail)

    def test_key_revoke_is_recorded(self):
        created = self.admin.post('/ui/keys', json={
            'label': 'k',
            'collections': [_server.DEFAULT_MANUAL_COLLECTION_ID],
            'types': ['ipv4-addr']}).get_json()
        r = self.admin.post(f"/ui/keys/{created['id']}/revoke")
        self.assertEqual(r.status_code, 200)
        row = self._one(action='key.revoke')
        self.assertEqual(row.target, created['prefix'])

    def test_user_lifecycle_is_recorded(self):
        user = self._mkuser('bob', 'viewer')
        self.assertEqual(self._one(action='user.create', target='bob').detail,
                         'role=viewer')
        r = self.admin.post(f"/ui/users/{user['id']}", json={'role': 'editor'})
        self.assertEqual(r.status_code, 200)
        self.assertIn('role', self._one(action='user.update', target='bob').detail)
        r = self.admin.post(f"/ui/users/{user['id']}/delete")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self._one(action='user.delete', target='bob').actor, 'root')

    def test_failed_request_writes_no_row(self):
        self._mkuser('ed', 'editor')
        editor = self._login('ed')
        r = editor.post('/feed/ingest', json={})     # 400: no stix_objects
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self._events(action='feed.ingest'), [])

    # --- window / filters --------------------------------------------------

    def test_limit_one_returns_only_the_newest(self):
        self._mkuser('ed', 'editor')
        editor = self._login('ed')
        for value in ('7.7.7.7', '8.8.8.8'):
            editor.post('/feed/ingest',
                        json={'stix_objects': [_manual_obj(value)]})
        body = self.admin.get('/ui/audit?limit=1').get_json()
        self.assertEqual(body['count'], 1)
        self.assertTrue(body['truncated'])
        self.assertEqual(body['events'][0]['action'], 'feed.ingest')

    def test_limit_is_validated(self):
        self.assertEqual(self.admin.get('/ui/audit?limit=0').status_code, 400)
        self.assertEqual(self.admin.get('/ui/audit?limit=abc').status_code, 400)

    def test_action_and_actor_filters(self):
        self._mkuser('ed', 'editor')
        editor = self._login('ed')
        editor.post('/feed/ingest', json={'stix_objects': [_manual_obj('7.7.7.7')]})
        self.admin.post('/ui/keys', json={
            'label': 'k', 'collections': [_server.DEFAULT_MANUAL_COLLECTION_ID],
            'types': ['ipv4-addr']})
        only_ingest = self.admin.get('/ui/audit?action=feed.ingest').get_json()
        self.assertTrue(only_ingest['events'])
        self.assertTrue(all(e['action'] == 'feed.ingest'
                            for e in only_ingest['events']))
        only_ed = self.admin.get('/ui/audit?actor=ed').get_json()
        self.assertTrue(only_ed['events'])
        self.assertTrue(all(e['actor'] == 'ed' for e in only_ed['events']))
        # An unknown value is an empty result, not an error.
        unknown = self.admin.get('/ui/audit?actor=nobody&action=no.such')
        self.assertEqual(unknown.status_code, 200)
        self.assertEqual(unknown.get_json()['events'], [])

    def test_newest_first_ordering(self):
        session = _server.create_session()
        try:
            session.query(_server.AuditEvent).delete()
            now = datetime.utcnow()
            for i, action in enumerate(
                    ('feed.ingest', 'feed.purge', 'feed.delete')):
                session.add(_server.AuditEvent(
                    id=f'id{i}', actor='root', action=action,
                    created=now - timedelta(seconds=i)))
            session.commit()
        finally:
            session.close()
        events = self.admin.get('/ui/audit').get_json()['events']
        self.assertEqual([e['action'] for e in events],
                         ['feed.ingest', 'feed.purge', 'feed.delete'])

    def test_audit_write_failure_does_not_break_the_request(self):
        # Best-effort: if appending the row blows up, the mutation still succeeds.
        self._mkuser('ed', 'editor')
        editor = self._login('ed')
        editor.post('/feed/ingest', json={'stix_objects': [_manual_obj('7.7.7.7')]})
        stix_id = _manual_obj('7.7.7.7')['id']
        original = _server.AuditEvent

        class _Boom:
            def __init__(self, *args, **kwargs):
                raise RuntimeError('audit store unavailable')

        _server.AuditEvent = _Boom
        try:
            r = editor.post(f'/objects/{stix_id}/revoke', json={'action': 'revoke'})
        finally:
            _server.AuditEvent = original
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.get_json()['revoked'])


class TestUiPageWiring(unittest.TestCase):
    """Static checks on intel-ui.html — the dashboard's page wiring.

    A dashboard page is defined in THREE places: the tab button, the
    ``<section>``, and the ``PAGE_NAMES`` allow-list. ``showPage()`` falls
    back to Feed for a name missing from ``PAGE_NAMES``, so a half-added page
    renders a visible tab that silently never opens. These assertions keep the
    three in step; there is no browser in CI to catch it at runtime.
    """

    def setUp(self):
        root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        with open(os.path.join(root, 'intel-ui.html'), encoding='utf-8') as fh:
            self.html = fh.read()

    def test_tabs_sections_and_page_names_agree(self):
        tabs = set(re.findall(
            r'<button class="tab[^"]*" data-page="([^"]+)"', self.html))
        sections = set(re.findall(r'<section data-page="([^"]+)"', self.html))
        raw = re.search(r'const PAGE_NAMES = \[([^\]]*)\]', self.html).group(1)
        names = {p.strip().strip('"\'') for p in raw.split(',') if p.strip()}

        self.assertTrue(tabs, 'no dashboard tabs found — did the markup change?')
        self.assertEqual(tabs, sections,
                         'every tab needs exactly one matching <section>')
        missing = tabs - names
        self.assertFalse(
            missing,
            f'{sorted(missing)} have a tab and a section but are missing from '
            f'PAGE_NAMES, so showPage() silently falls back to Feed')
        self.assertEqual(names, tabs, 'PAGE_NAMES lists a page that has no tab')


class _FakeJwkClient:
    """Stand-in for jwt.PyJWKClient: returns the local public key for any token."""

    def __init__(self, public_key):
        self._pub = public_key

    def get_signing_key_from_jwt(self, token):
        import types
        key = types.SimpleNamespace()
        key.key = self._pub
        return key


class TestSchemaDdl(unittest.TestCase):
    """The schema is auto-created from the ORM models, and the DDL a DBA-managed
    deployment is told to run (IMPLEMENTATION_GUIDE.md) is generated from those
    same models by tools/dump_schema.py. A table/column/index added, renamed or
    dropped without updating that script (and the guide) fails here — which is
    how api_keys/ui_users were missing from the docs despite being live tables.
    """

    EXPECTED_TABLES = {'stix_objects', 'subscriptions', 'api_keys',
                       'ui_users', 'puller_state', 'audit_events',
                       'benign_values'}
    POLL_INDEXES = {'ix_stix_objects_modified_id',
                    'ix_stix_objects_coll_modified_id'}

    def _dump_schema(self):
        import importlib.util
        path = os.path.join(os.path.dirname(os.path.dirname(__file__)),
                            'tools', 'dump_schema.py')
        spec = importlib.util.spec_from_file_location('dump_schema', path)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module

    def test_models_define_every_expected_table(self):
        self.assertEqual(set(_server.Base.metadata.tables), self.EXPECTED_TABLES,
                         'a model was added/renamed/dropped — update '
                         'tools/dump_schema.py and the guide DDL to match')

    def test_later_release_columns_are_present(self):
        # api_keys.max_objects and ui_users.role arrived after the tables first
        # shipped; both are load-bearing for the per-key ceiling and RBAC.
        self.assertIn('max_objects',
                      _server.Base.metadata.tables['api_keys'].columns)
        self.assertIn('role', _server.Base.metadata.tables['ui_users'].columns)

    def test_generated_ddl_covers_tables_and_poll_indexes(self):
        # PostgreSQL is the deployment target (docker compose); the DDL a
        # DBA runs is the Postgres dialect from tools/dump_schema.py.
        ddl = self._dump_schema()._ddl_for(_server.Base.metadata, 'postgresql')
        for table in self.EXPECTED_TABLES:
            self.assertIn(f'CREATE TABLE {table} ', ddl)
        for index in self.POLL_INDEXES:
            self.assertIn(index, ddl)

    def test_check_mode_passes_on_current_models(self):
        self.assertEqual(self._dump_schema()._check(_server.Base.metadata), 0)


class TestBenignFilter(unittest.TestCase):
    """Known-legitimate filter — keeps legit values OUT of the Vision One feed.

    A value named in ``intel_filter.benign_values`` / ``benign_domains`` /
    ``benign_cidrs`` (config.yaml) or in the ``benign_values`` table (dashboard,
    via ``/ui/benign`` or the one-click ``/objects/<id>/exclude``) is WITHHELD
    from the feed — Vision One never sees it, so it cannot block real traffic on
    it. Manual intel is never gated.
    """

    ADMIN = ('root', 'root-password-123')
    PW = 'password-12345'

    def setUp(self):
        self.client = app.test_client()
        _reset_db()
        app.config['TAXII_AUTH'] = {'username': 'u', 'password': 'p'}
        app.config['TAXII_ADMIN_AUTH'] = {'username': 'u', 'password': 'p'}
        self.auth = {'X-Taxii-Username': 'u', 'X-Taxii-Password': 'p'}

        # Admin session — the write principal for the dashboard endpoints.
        self._saved_creds = (_server._ui_username, _server._ui_password)
        _server._ui_username, _server._ui_password = self.ADMIN
        self.admin = app.test_client()
        r = self.admin.post('/ui/login', json={
            'username': self.ADMIN[0], 'password': self.ADMIN[1]})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))

        # Snapshot the static benign lists a test may monkeypatch.
        self._saved_lists = (
            _server.FILTER_BENIGN_VALUES,
            _server.FILTER_BENIGN_DOMAINS,
            _server.FILTER_BENIGN_CIDRS,
        )

    def tearDown(self):
        _server._ui_username, _server._ui_password = self._saved_creds
        (_server.FILTER_BENIGN_VALUES,
         _server.FILTER_BENIGN_DOMAINS,
         _server.FILTER_BENIGN_CIDRS) = self._saved_lists

    # --- helpers -----------------------------------------------------------

    def _community_ip(self, value):
        ingest_objects([
            {'id': f'ipv4-addr--{value.replace(".", "-")}', 'type': 'ipv4-addr',
             'object': {'ipv4-addr': {'value': value}},
             'labels': ['otx'], 'confidence': 70}
        ], mode='merge', source='otx')

    def _community_domain(self, value):
        ingest_objects([
            {'id': f'domain-name--{value.replace(".", "-")}', 'type': 'domain-name',
             'object': {'domain-name': {'value': value}},
             'labels': ['otx'], 'confidence': 70}
        ], mode='merge', source='otx')

    def _taxii_values(self):
        r = self.client.get(
            '/taxii2/collections/threat_intel/objects/', headers=self.auth)
        self.assertEqual(r.status_code, 200)
        content = r.get_json()['content']['content']['objects']
        return {served_value(o) for o in content if served_value(o)}

    def _object_row(self, value):
        return next(
            o for o in self.client.get('/objects', headers=self.auth)
            .get_json()['objects'] if o['value'] == value)

    def _actions(self):
        session = _server.create_session()
        try:
            return [e.action for e in session.query(_server.AuditEvent).all()]
        finally:
            session.close()

    # --- static benign lists (config.yaml) ---------------------------------

    def test_config_benign_value_withholds(self):
        _server.FILTER_BENIGN_VALUES = {'8.8.8.8'}
        self._community_ip('8.8.8.8')          # public -> would be served
        self.assertNotIn('8.8.8.8', self._taxii_values())

    def test_config_benign_domain_covers_subdomains(self):
        _server.FILTER_BENIGN_DOMAINS = {'example.com'}
        self._community_domain('login.example.com')
        self.assertNotIn('login.example.com', self._taxii_values())

    def test_config_benign_cidr_withholds(self):
        _server.FILTER_BENIGN_CIDRS = _server._parse_benign_cidrs(['13.107.0.0/16'])
        self._community_ip('13.107.5.5')
        self.assertNotIn('13.107.5.5', self._taxii_values())

    def test_unrelated_value_is_still_served(self):
        """A benign entry must not withhold anything but its own value."""
        _server.FILTER_BENIGN_VALUES = {'8.8.8.8'}
        self._community_ip('8.8.8.8')
        self._community_ip('9.9.9.9')
        served = self._taxii_values()
        self.assertNotIn('8.8.8.8', served)
        self.assertIn('9.9.9.9', served)

    def test_manual_intel_is_never_filtered(self):
        _server.FILTER_BENIGN_DOMAINS = {'example.com'}
        ingest_objects([
            {'id': 'domain-name--login-example-com', 'type': 'domain-name',
             'object': {'domain-name': {'value': 'login.example.com'}},
             'labels': ['manual'], 'confidence': 90}
        ], mode='replace', source='manual')
        self.assertIn('login.example.com', self._taxii_values())

    # --- dynamic benign entries (dashboard) --------------------------------

    def test_create_benign_withholds_a_served_object(self):
        self._community_ip('8.8.8.8')
        self.assertIn('8.8.8.8', self._taxii_values())
        r = self.admin.post('/ui/benign', json={
            'value': '8.8.8.8', 'kind': 'exact', 'note': 'public DNS'})
        self.assertEqual(r.status_code, 201)
        body = r.get_json()
        self.assertTrue(body['changed'])
        self.assertEqual(body['value'], '8.8.8.8')
        self.assertEqual(body['kind'], 'exact')
        self.assertNotIn('8.8.8.8', self._taxii_values())

        # Removing it serves the object again.
        self.assertEqual(
            self.admin.post(f"/ui/benign/{body['id']}/delete").status_code, 200)
        self.assertIn('8.8.8.8', self._taxii_values())

    def test_exclude_object_route_defaults_to_exact_for_ip(self):
        self._community_ip('8.8.4.4')
        r = self.admin.post('/objects/ipv4-addr--8-8-4-4/exclude', json={})
        self.assertEqual(r.status_code, 201, r.get_data(as_text=True))
        body = r.get_json()
        self.assertEqual(body['kind'], 'exact')
        self.assertEqual(body['value'], '8.8.4.4')
        self.assertNotIn('8.8.4.4', self._taxii_values())

    def test_exclude_object_route_defaults_to_domain_for_domain(self):
        self._community_domain('www.example.com')
        self.assertIn('www.example.com', self._taxii_values())
        r = self.admin.post(
            '/objects/domain-name--www-example-com/exclude', json={})
        self.assertEqual(r.status_code, 201, r.get_data(as_text=True))
        self.assertEqual(r.get_json()['kind'], 'domain')
        self.assertNotIn('www.example.com', self._taxii_values())

    def test_exclude_object_unknown_404(self):
        r = self.admin.post('/objects/ipv4-addr--0-0-0-0/exclude', json={})
        self.assertEqual(r.status_code, 404)

    def test_objects_endpoint_reports_the_benign_gate(self):
        self._community_ip('8.8.8.8')
        before = self._object_row('8.8.8.8')
        self.assertFalse(before['gated'])
        self.assertIsNone(before['gate_reason'])

        self.admin.post('/objects/ipv4-addr--8-8-8-8/exclude', json={})
        after = self._object_row('8.8.8.8')
        self.assertTrue(after['gated'])
        self.assertEqual(after['gate_reason'], 'benign')
        self.assertEqual(after['gate_reason_label'],
                         'Known legitimate (kept out of the feed)')

    def test_create_is_idempotent(self):
        first = self.admin.post('/ui/benign', json={
            'value': 'a.example.com', 'kind': 'domain'})
        self.assertEqual(first.status_code, 201)
        second = self.admin.post('/ui/benign', json={
            'value': 'a.example.com', 'kind': 'domain'})
        self.assertEqual(second.status_code, 200)
        self.assertFalse(second.get_json()['changed'])
        self.assertEqual(self.admin.get('/ui/benign').get_json()['count'], 1)

    def test_list_includes_config_rules_and_delete_unknown_404(self):
        self.admin.post('/ui/benign', json={'value': '8.8.8.8', 'kind': 'exact'})
        body = self.admin.get('/ui/benign').get_json()
        self.assertEqual(body['count'], 1)
        self.assertIn('exact', body['kinds'])
        self.assertIn('config', body)
        self.assertIn('benign_values', body['config'])
        self.assertEqual(
            self.admin.post('/ui/benign/nope/delete').status_code, 404)

    # --- validation --------------------------------------------------------

    def test_create_rejects_bad_kind_and_values(self):
        for payload in (
            {'value': 'x.example.com', 'kind': 'nope'},
            {'value': '', 'kind': 'exact'},
            {'value': 'not a domain', 'kind': 'domain'},
            {'value': '999.999.1.1/24', 'kind': 'cidr'},
        ):
            r = self.admin.post('/ui/benign', json=payload)
            self.assertEqual(r.status_code, 400, payload)

    # --- access control ----------------------------------------------------

    def test_writes_require_a_write_role(self):
        self.admin.post('/ui/users', json={
            'username': 'vi', 'password': self.PW, 'role': 'viewer'})
        viewer = app.test_client()
        self.assertEqual(viewer.post('/ui/login', json={
            'username': 'vi', 'password': self.PW}).status_code, 200)

        # A viewer may read the list but not change it.
        self.assertEqual(viewer.get('/ui/benign').status_code, 200)
        self.assertEqual(viewer.post('/ui/benign', json={
            'value': 'x.example.com', 'kind': 'domain'}).status_code, 401)
        self.assertEqual(
            viewer.post('/objects/ipv4-addr--1-2-3-4/exclude', json={}).status_code,
            401)
        self.assertEqual(app.test_client().post(
            '/ui/benign', json={'value': 'x'}).status_code, 401)

    # --- audit -------------------------------------------------------------

    def test_benign_changes_are_audited(self):
        created = self.admin.post('/ui/benign', json={
            'value': '8.8.8.8', 'kind': 'exact'}).get_json()
        self.admin.post(f"/ui/benign/{created['id']}/delete")
        actions = self._actions()
        self.assertIn('benign.create', actions)
        self.assertIn('benign.delete', actions)


if __name__ == '__main__':
    unittest.main(verbosity=2)
